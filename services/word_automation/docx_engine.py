"""Генерация и нормализация DOCX по пресету ``StylePreset`` (python-docx).

Класс :class:`DocxFormatter` решает две задачи:

* :meth:`DocxFormatter.create_document` — собирает чистый ``.docx`` из
  древовидной/списочной структуры контента строго по правилам пресета;
* :meth:`DocxFormatter.normalize_document` — открывает существующий ``.docx``,
  вычищает «грязное» оформление (серии пробелов/табуляций, пустые параграфы
  вместо отступов, сбитые инлайн-шрифты и кегли) и принудительно накатывает
  пресет на секции, стили, параграфы и таблицы.

Единицы пресета: миллиметры (поля, отступ первой строки) и пункты (кегли,
интервалы до/после). Пересчёт в EMU выполняют ``docx.shared.Mm``/``Pt``.
"""

from __future__ import annotations

import os
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Sequence

import docx
from docx.document import Document as DocumentObject
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_BREAK,
    WD_LINE_SPACING,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run

try:  # импорт внутри пакета
    from .preset_schema import (
        Alignment,
        CaptionPlacement,
        CellVerticalAlignment,
        CleanupMode,
        CleanupPolicy,
        HeadingStyle,
        LineSpacingRule,
        ListStyle,
        POINT_BASED_RULES,
        PageNumberingStyle,
        PresetError,
        SPECIAL_SECTION_KINDS,
        SignatureLineStyle,
        SpacerMode,
        SpecialSectionsStyle,
        StylePreset,
        TableStyle,
        TextCase,
        TitlePageStyle,
        _log,
        resolve_preset,
    )
except ImportError:  # pragma: no cover - запуск модуля вне пакета
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from preset_schema import (  # type: ignore[no-redef]
        Alignment,
        CaptionPlacement,
        CellVerticalAlignment,
        CleanupMode,
        CleanupPolicy,
        HeadingStyle,
        LineSpacingRule,
        ListStyle,
        POINT_BASED_RULES,
        PageNumberingStyle,
        PresetError,
        SPECIAL_SECTION_KINDS,
        SignatureLineStyle,
        SpacerMode,
        SpecialSectionsStyle,
        StylePreset,
        TableStyle,
        TextCase,
        TitlePageStyle,
        _log,
        resolve_preset,
    )


_APP_DIR = Path(__file__).resolve().parents[2]
OUTPUT_DIRNAME = os.path.join("generated_files", "word_automation")

#: Каталог по умолчанию для создаваемых документов.
from app_paths import USER_DATA_DIR

DEFAULT_OUTPUT_DIR = USER_DATA_DIR / OUTPUT_DIRNAME


class DocxStructureError(PresetError):
    """Передана непонятная или пустая структура контента."""


class DocxFileError(PresetError):
    """Ошибка чтения/записи файла документа."""


# --------------------------------------------------------------------------------------
# Соответствия пресет -> перечисления Word
# --------------------------------------------------------------------------------------

ALIGNMENT_MAP: dict[Alignment, Any] = {
    Alignment.LEFT: WD_ALIGN_PARAGRAPH.LEFT,
    Alignment.CENTER: WD_ALIGN_PARAGRAPH.CENTER,
    Alignment.RIGHT: WD_ALIGN_PARAGRAPH.RIGHT,
    Alignment.JUSTIFY: WD_ALIGN_PARAGRAPH.JUSTIFY,
}

LINE_RULE_MAP: dict[LineSpacingRule, Any] = {
    LineSpacingRule.SINGLE: WD_LINE_SPACING.SINGLE,
    LineSpacingRule.ONE_POINT_FIVE: WD_LINE_SPACING.ONE_POINT_FIVE,
    LineSpacingRule.DOUBLE: WD_LINE_SPACING.DOUBLE,
    LineSpacingRule.MULTIPLE: WD_LINE_SPACING.MULTIPLE,
    LineSpacingRule.EXACTLY: WD_LINE_SPACING.EXACTLY,
    LineSpacingRule.AT_LEAST: WD_LINE_SPACING.AT_LEAST,
}

#: Границы таблицы в порядке записи ``w:tblBorders``.
TABLE_BORDER_EDGES = ("top", "left", "bottom", "right", "insideH", "insideV")

#: Выравнивание таблицы на странице (``w:jc`` в ``w:tblPr``).
TABLE_ALIGNMENT_MAP: dict[Alignment, Any] = {
    Alignment.LEFT: WD_TABLE_ALIGNMENT.LEFT,
    Alignment.CENTER: WD_TABLE_ALIGNMENT.CENTER,
    Alignment.RIGHT: WD_TABLE_ALIGNMENT.RIGHT,
    Alignment.JUSTIFY: WD_TABLE_ALIGNMENT.CENTER,
}

#: Вертикальное выравнивание текста в ячейке (``w:vAlign``).
CELL_VALIGN_MAP: dict[CellVerticalAlignment, Any] = {
    CellVerticalAlignment.TOP: WD_ALIGN_VERTICAL.TOP,
    CellVerticalAlignment.CENTER: WD_ALIGN_VERTICAL.CENTER,
    CellVerticalAlignment.BOTTOM: WD_ALIGN_VERTICAL.BOTTOM,
}

#: Стандартные размеры листа A4 (мм) — используются, если в секции нет размеров.
A4_SHORT_MM, A4_LONG_MM = 210.0, 297.0

#: Максимальная глубина вложенности списка (дальше отступ перестаёт расти).
MAX_LIST_LEVEL = 4

#: Приблизительная ширина символа Times New Roman как доля кегля (для эвристик).
_AVG_CHAR_WIDTH_RATIO = 0.5
#: Ёмкость строки по умолчанию, если ширина текста получилась неположительной.
_FALLBACK_LINE_CAPACITY = 70

#: Подпись вида «Таблица 1 — ...», «Рисунок 2.1 – ...», «Приложение А ...».
_CAPTION_RE = re.compile(
    r"^(таблиц|рисунк|рис\.|figure|table|приложени|appendix|схема|чертёж|чертеж)"
    r"[а-яa-z]*\s*[№\dА-ЯA-Z.]",
    re.IGNORECASE,
)
#: Ведущая нумерация элемента списка: «1.», «1)», «1.2.», «а)».
_NUMBERED_LEAD_RE = re.compile(r"^(?:\d+(?:\.\d+)*[\.\)]|[а-яa-z][\.\)])\s+\S")
#: Уже проставленный маркер пункта списка (номер или маркер) — чтобы не дублить.
_LIST_MARKER_RE = re.compile(
    r"^(?:\d+(?:\.\d+)*[\.\)]|[а-яa-z][\.\)]|[–—•·\-\*o])\s+\S", re.IGNORECASE
)
#: «Мягкое» распознавание уже существующего маркера: тире/буллет/номер/буква,
#: пробел после маркера не обязателен (иначе «-текст» и «а)текст» дублировались).
_LIST_MARKER_PREFIX_RE = re.compile(
    r"^(?:[–—•·\*]\s*|-\s*|[оo]\s+|\d{1,3}(?:\.\d+)*[\.\)]\s*"
    r"|[а-яёa-zА-ЯЁA-Z][\.\)]\s*)",
    re.IGNORECASE,
)
#: Разделение строки на «маркер + текст»: маркер — буллет, номер или буква.
#: Используется, чтобы снять ИСХОДНЫЙ маркер перед наложением маркера пресета
#: (иначе получалось «- - пункт» или «1. 1. пункт»).
_LEADING_MARKER_RE = re.compile(
    r"^\s*(?P<marker>\d{1,3}(?:\.\d+)*[.\)]|[а-яёa-zА-ЯЁA-Z][.\)]|[–—•·\*\-])"
    r"(?P<sep>\s+)?(?P<rest>\S.*)$",
    re.UNICODE,
)
#: Числовой маркер («1.», «2)», «1.2.») — признак НАСТОЯЩЕЙ нумерации исходника.
_NUMERIC_MARKER_RE = re.compile(r"^\d{1,3}(?:\.\d+)*[.\)]$")
#: Нумерованный заголовок раздела: «1 ГЛАВА…», «1. АНАЛИЗ…», «1.1 Описание…».
_SECTION_NUMBER_RE = re.compile(
    r"^(?P<number>\d{1,3}(?:\.\d{1,3})*)(?P<delim>\.|\s)\s*(?P<title>\S.*)$",
    re.UNICODE,
)
#: Максимальное число слов в строке, чтобы считать её заголовком раздела.
_MAX_HEADING_WORDS = 12


#: Соответствие канонического заполнителя табуляции перечислению python-docx.
_TAB_LEADER_MAP: dict[str, Any] = {
    "dot": WD_TAB_LEADER.DOTS,
    "hyphen": WD_TAB_LEADER.DASHES,
    "underscore": WD_TAB_LEADER.LINES,
    "none": WD_TAB_LEADER.SPACES,
}

#: Структурные разделы, которые должны начинаться со строго новой страницы
#: (используется, если пресет не задал собственный список).
_DEFAULT_STRUCTURAL_TITLES = (
    "реферат", "содержание", "оглавление", "введение", "заключение",
    "список", "перечень", "приложение", "обозначен", "сокращен",
    "определения", "термины", "аннотация", "выводы",
)

#: Символы, считающиеся пробельными при очистке «грязного» текста.
_WHITESPACE_CHARS = " \t\u00a0\u2000\u2001\u2002\u2003\u2007\u202f\u205f\u3000"

_MULTISPACE_RE = re.compile(r"[ \t\u00a0\u2000-\u200b\u202f\u205f\u3000]{2,}")
_LEADING_WS_RE = re.compile(r"^[ \t\u00a0]+")
_UNDERLINE_GAP_RE = re.compile(r"_{3,}")
#: Inline-разметка: **жирный**, *курсив*, __жирный__, _курсив_.
_INLINE_MARK_RE = re.compile(
    r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|__[^_]+__|\*[^*\n]+\*|_[^_\n]+_)"
)

#: Порог «костыльного» смещения текста пробелами (ТЗ: более 3 подряд).
SPACE_SHIFT_THRESHOLD = 3
#: Максимум пустых параграфов, переводимых в ``space_before``.
MAX_EMPTY_PARAGRAPHS_AS_SPACE = 4


# --------------------------------------------------------------------------------------
# Модель контент-блока
# --------------------------------------------------------------------------------------


@dataclass(frozen=True)
class RunSpec:
    """Фрагмент текста с оформлением (для inline-разметки и явных ``runs``)."""

    text: str = ""
    bold: bool | None = None
    italic: bool | None = None
    underline: bool | None = None


@dataclass(frozen=True)
class ContentBlock:
    """Нормализованный блок контента (заголовок, абзац, список, таблица, подпись)."""

    kind: str
    level: int = 0
    text: str = ""
    runs: tuple[RunSpec, ...] = ()
    items: tuple[str, ...] = ()
    ordered: bool = False
    header: tuple[str, ...] = ()
    rows: tuple[tuple[str, ...], ...] = ()
    caption: str = ""
    caption_kind: str = "table"
    caption_number: Any = None
    alignment: Alignment | None = None
    bold: bool | None = None
    italic: bool | None = None
    first_line_indent_mm: float | None = None
    #: Сырые данные блока (для титульного листа и спецразделов).
    payload: dict[str, Any] = field(default_factory=dict)


#: Синонимы типов блоков -> канонический ``kind``.
BLOCK_TYPE_ALIASES: dict[str, str] = {
    "h1": "heading", "h2": "heading", "h3": "heading",
    "h4": "heading", "h5": "heading", "h6": "heading",
    "heading1": "heading", "heading2": "heading", "heading3": "heading",
    "heading_1": "heading", "heading_2": "heading", "heading_3": "heading",
    "heading": "heading", "title": "heading", "section": "heading",
    "chapter": "heading", "subsection": "heading", "subsubsection": "heading",
    "p": "paragraph", "paragraph": "paragraph", "text": "paragraph",
    "para": "paragraph", "body": "paragraph", "quote": "paragraph",
    "table": "table", "tbl": "table", "grid": "table",
    "list": "list", "ul": "list", "ol": "list", "bullet": "list",
    "bullets": "list", "numbered": "list", "ordered": "list",
    "caption": "caption", "figure_caption": "caption", "table_caption": "caption",
    "page_break": "page_break", "pagebreak": "page_break",
    "break": "page_break", "new_page": "page_break",
    # Титульный лист и специальные разделы (реферат/задание/содержание).
    "titlepage": "title_page", "title_page": "title_page", "cover": "title_page",
    "титульный": "title_page", "титульник": "title_page", "титул": "title_page",
    "task": "task", "zadanie": "task", "задание": "task",
    "referat": "referat", "abstract": "referat", "реферат": "referat",
    "content": "content", "toc": "content", "содержание": "content",
    "оглавление": "content", "contents": "content",
}

#: Типы блоков, которые рендерятся как специальные секции документа.
SPECIAL_BLOCK_KINDS: frozenset[str] = frozenset(
    {"title_page", "task", "referat", "content"}
)

#: Уровень заголовка по исходному имени типа.
_TYPE_TO_LEVEL: dict[str, int] = {
    "h1": 1, "h2": 2, "h3": 3, "h4": 4, "h5": 5, "h6": 6,
    "heading1": 1, "heading2": 2, "heading3": 3,
    "heading_1": 1, "heading_2": 2, "heading_3": 3,
    "title": 1, "section": 1, "chapter": 1,
    "subsection": 2, "subsubsection": 3,
}

SUPPORTED_BLOCK_TYPES = (
    "h1, h2, h3 (заголовки); paragraph/p (абзац); list/ul/ol (список); "
    "table (таблица); caption (подпись); page_break (разрыв страницы)"
)


def _as_text(value: Any) -> str:
    """Аккуратно приводит значение к тексту (списки склеиваются через пробел)."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        return str(value)
    if isinstance(value, (list, tuple)):
        return " ".join(_as_text(item) for item in value if item is not None)
    return str(value)


def _first_present(data: dict[str, Any], keys: Sequence[str]) -> Any:
    """Первое непустое значение среди перечисленных ключей словаря."""
    for key in keys:
        if key in data and data[key] not in (None, ""):
            return data[key]
    return None


def _as_bool(value: Any, default: bool | None = None) -> bool | None:
    """Приводит значение к ``bool``/``None`` (``None`` — «не переопределять»)."""
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    text = str(value).strip().lower()
    if text in ("1", "true", "yes", "y", "да", "on", "вкл"):
        return True
    if text in ("0", "false", "no", "n", "нет", "off", "выкл"):
        return False
    return default


def _as_alignment(value: Any) -> Alignment | None:
    """Разбирает выравнивание из строки (в том числе русской)."""
    if value is None:
        return None
    if isinstance(value, Alignment):
        return value
    text = str(value).strip().lower()
    if not text:
        return None
    mapping = {
        "left": Alignment.LEFT, "l": Alignment.LEFT,
        "по левому": Alignment.LEFT, "слева": Alignment.LEFT,
        "левое": Alignment.LEFT,
        "center": Alignment.CENTER, "centre": Alignment.CENTER,
        "c": Alignment.CENTER, "по центру": Alignment.CENTER,
        "центр": Alignment.CENTER, "центральное": Alignment.CENTER,
        "right": Alignment.RIGHT, "r": Alignment.RIGHT,
        "по правому": Alignment.RIGHT, "справа": Alignment.RIGHT,
        "правое": Alignment.RIGHT,
        "justify": Alignment.JUSTIFY, "j": Alignment.JUSTIFY,
        "по ширине": Alignment.JUSTIFY, "justified": Alignment.JUSTIFY,
        "ширина": Alignment.JUSTIFY,
    }
    result = mapping.get(text)
    if result is None:
        try:
            result = Alignment(text)
        except ValueError:
            _log(f"word_automation: неизвестное выравнивание '{value}' — пропущено")
            return None
    return result


def _cell_text(value: Any) -> str:
    """Текст ячейки таблицы: ``None`` -> пустая строка, списки -> через '; '."""
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        return "; ".join(_as_text(item) for item in value if item is not None)
    return _as_text(value)


def _pad_rows(
    rows: tuple[tuple[str, ...], ...], columns: int
) -> tuple[tuple[str, ...], ...]:
    """Дополняет строки до ``columns`` ячеек (защита от «рваных» таблиц)."""
    columns = max(1, columns)
    padded: list[tuple[str, ...]] = []
    for row in rows:
        if len(row) < columns:
            padded.append(row + ("",) * (columns - len(row)))
        else:
            padded.append(row[:columns])
    return tuple(padded)


def _normalize_rows(raw: Any, columns: int) -> tuple[tuple[str, ...], ...]:
    """Приводит строки таблицы к кортежам одинаковой ширины.

    Поддерживаются формы: список списков, список словарей, словарь колонок
    ``{"col": [...]}`` и плоский список (одно значение = одна строка из 1 ячейки).
    """
    if raw is None:
        return ()
    if isinstance(raw, dict):
        keys = list(raw.keys())
        columns = max(columns, len(keys))
        height = 0
        for value in raw.values():
            if isinstance(value, (list, tuple)):
                height = max(height, len(value))
        result: list[tuple[str, ...]] = []
        for index in range(height):
            row: list[str] = []
            for key in keys:
                values = raw[key]
                if isinstance(values, (list, tuple)):
                    row.append(_cell_text(values[index]) if index < len(values) else "")
                else:
                    row.append(_cell_text(values) if index == 0 else "")
            result.append(tuple(row))
        return _pad_rows(tuple(result), columns)
    if isinstance(raw, (list, tuple)):
        rows: list[tuple[str, ...]] = []
        for item in raw:
            if isinstance(item, (list, tuple)):
                rows.append(tuple(_cell_text(cell) for cell in item))
            elif isinstance(item, dict):
                rows.append(tuple(_cell_text(value) for value in item.values()))
            else:
                rows.append((_cell_text(item),))
        width = max([columns] + [len(row) for row in rows])
        return _pad_rows(tuple(rows), width)
    raise DocxStructureError(
        "Поле rows таблицы должно быть списком списков/словарей, "
        f"получено: {type(raw).__name__}"
    )


def parse_inline_runs(text: str) -> tuple[RunSpec, ...]:
    """Разбирает inline-разметку ``**жирный**`` / ``*курсив*`` на фрагменты.

    Одиночные маркеры без пары остаются обычным текстом, поэтому разметка
    безопасна для текста вроде ``2 * 3`` или ``И.О. _Фамилия``.
    """
    source = _as_text(text)
    if not source or not _INLINE_MARK_RE.search(source):
        return (RunSpec(text=source),)
    specs: list[RunSpec] = []
    for part in _INLINE_MARK_RE.split(source):
        if not part:
            continue
        if part.startswith("***") and part.endswith("***") and len(part) > 6:
            specs.append(RunSpec(part[3:-3], bold=True, italic=True))
        elif part.startswith("**") and part.endswith("**") and len(part) > 4:
            specs.append(RunSpec(part[2:-2], bold=True))
        elif part.startswith("__") and part.endswith("__") and len(part) > 4:
            specs.append(RunSpec(part[2:-2], bold=True))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            specs.append(RunSpec(part[1:-1], italic=True))
        elif part.startswith("_") and part.endswith("_") and len(part) > 2:
            specs.append(RunSpec(part[1:-1], italic=True))
        else:
            specs.append(RunSpec(part))
    return tuple(specs) or (RunSpec(text=source),)



# --------------------------------------------------------------------------------------
# Разбор структуры контента
# --------------------------------------------------------------------------------------


def _resolve_block_kind(data: dict[str, Any]) -> tuple[str, int]:
    """Определяет канонический тип блока и уровень заголовка.

    Возвращает ``(kind, level)``; ``level`` значим только для ``kind='heading'``.
    """
    raw_type = _first_present(data, ("type", "kind", "block", "block_type", "style"))
    if raw_type is None:
        # Тип не указан — угадываем по составу полей.
        if _first_present(data, ("rows", "data", "header", "headers", "columns")) is not None:
            return "table", 0
        if _first_present(data, ("items", "entries")) is not None:
            return "list", 0
        return "paragraph", 0

    text = str(raw_type).strip().lower().replace("-", "").replace(" ", "").replace("_", "")
    key = str(raw_type).strip().lower().replace("-", "").replace(" ", "")
    kind = BLOCK_TYPE_ALIASES.get(text) or BLOCK_TYPE_ALIASES.get(key)
    if kind is None:
        raise DocxStructureError(
            f"Неизвестный тип блока '{raw_type}'. Поддерживаются: {SUPPORTED_BLOCK_TYPES}"
        )
    level = _TYPE_TO_LEVEL.get(text) or _TYPE_TO_LEVEL.get(key) or 0
    return kind, level


def _explicit_runs(data: dict[str, Any]) -> tuple[RunSpec, ...]:
    """Явные фрагменты ``runs`` из блока (если заданы)."""
    raw = data.get("runs")
    if not isinstance(raw, (list, tuple)) or not raw:
        return ()
    specs: list[RunSpec] = []
    for item in raw:
        if isinstance(item, str):
            specs.append(RunSpec(text=item))
        elif isinstance(item, dict):
            specs.append(
                RunSpec(
                    text=_as_text(_first_present(item, ("text", "content", "value"))),
                    bold=_as_bool(item.get("bold")),
                    italic=_as_bool(item.get("italic")),
                    underline=_as_bool(item.get("underline")),
                )
            )
        else:
            specs.append(RunSpec(text=_as_text(item)))
    return tuple(spec for spec in specs if spec.text)


def _build_table_block(data: dict[str, Any]) -> ContentBlock:
    """Собирает блок таблицы из словаря, выравнивая ширину header/rows."""
    header_raw = _first_present(data, ("header", "headers", "columns", "thead"))
    header: tuple[str, ...] = ()
    if isinstance(header_raw, (list, tuple)):
        header = tuple(_cell_text(cell) for cell in header_raw)
    elif header_raw not in (None, ""):
        header = (_cell_text(header_raw),)
    rows = _normalize_rows(
        _first_present(data, ("rows", "data", "body", "tbody")), len(header)
    )
    if not header and not rows:
        raise DocxStructureError("Блок table пуст: передайте header и/или rows.")
    if header and rows and len(header) != len(rows[0]):
        width = max(len(header), len(rows[0]))
        header = _pad_rows((header,), width)[0]
        rows = _pad_rows(rows, width)
    return ContentBlock(
        kind="table",
        header=header,
        rows=rows,
        caption=_as_text(_first_present(data, ("caption", "title"))),
        caption_kind=str(
            _first_present(data, ("caption_type", "caption_kind")) or "table"
        ).strip().lower(),
        caption_number=_first_present(data, ("number", "caption_number", "n")),
        alignment=_as_alignment(_first_present(data, ("alignment", "align"))),
    )


def _build_list_block(data: dict[str, Any], kind_hint: str) -> ContentBlock:
    """Собирает блок списка из словаря (items либо текст по строкам)."""
    items_raw = _first_present(data, ("items", "entries", "values", "points"))
    if items_raw is None:
        text_value = _as_text(_first_present(data, ("text", "content", "value")))
        items_raw = [line for line in text_value.split("\n") if line.strip()]
    if isinstance(items_raw, (list, tuple)):
        items = tuple(_as_text(item).strip() for item in items_raw)
    else:
        items = (_as_text(items_raw).strip(),)
    items = tuple(item for item in items if item)
    if not items:
        raise DocxStructureError("Блок list пуст: передайте непустые items.")
    ordered = _as_bool(_first_present(data, ("ordered", "numbered", "numeric")))
    if ordered is None:
        ordered = kind_hint in ("ol", "ordered", "numbered")
    return ContentBlock(
        kind="list",
        items=items,
        ordered=bool(ordered),
        alignment=_as_alignment(_first_present(data, ("alignment", "align"))),
    )



def _build_caption_block(data: dict[str, Any]) -> ContentBlock:
    """Собирает блок подписи (рисунок или таблица)."""
    caption_kind = str(
        _first_present(data, ("caption_type", "caption_kind", "object")) or "figure"
    ).strip().lower()
    return ContentBlock(
        kind="caption",
        caption=_as_text(
            _first_present(data, ("text", "content", "value", "caption"))
        ),
        caption_kind=caption_kind,
        caption_number=_first_present(data, ("number", "caption_number", "n")),
        alignment=_as_alignment(_first_present(data, ("alignment", "align"))),
    )



def _build_special_block(data: dict[str, Any], kind: str) -> ContentBlock:
    """Собирает блок титульного листа или специального раздела.

    ``title_page`` хранит сырой словарь (шапка/заголовок/подписи/подвал).
    ``task``/``referat``/``content`` — заголовок раздела, абзацы тела и пункты
    списка (маркеры не переформатируются, чтобы сохранить дефисы и буквы).
    """
    payload = dict(data)
    if kind == "title_page":
        return ContentBlock(kind="title_page", payload=payload)

    title = _as_text(_first_present(data, ("title", "heading", "name", "text")))
    body: list[str] = []
    raw_body = _first_present(data, ("body", "paragraphs", "lines", "value", "content"))
    if isinstance(raw_body, (list, tuple)):
        body = [_as_text(item).strip() for item in raw_body if _as_text(item).strip()]
    elif raw_body not in (None, ""):
        body = [line.strip() for line in _as_text(raw_body).split("\n") if line.strip()]

    items: list[str] = []
    raw_items = _first_present(data, ("items", "entries", "points", "tasks", "list"))
    if isinstance(raw_items, (list, tuple)):
        items = [_as_text(item).strip() for item in raw_items if _as_text(item).strip()]
    elif raw_items not in (None, ""):
        items = [line.strip() for line in _as_text(raw_items).split("\n") if line.strip()]

    payload["_title"] = title
    payload["_body"] = tuple(body)
    payload["_items"] = tuple(items)
    payload["_ordered"] = bool(
        _as_bool(_first_present(data, ("ordered", "numbered", "numeric")))
    )
    return ContentBlock(
        kind=kind, text=title, items=tuple(items), payload=payload
    )


def _block_from_dict(data: dict[str, Any], inherited_level: int = 0) -> list[ContentBlock]:
    """Собирает блок (и его вложенных детей) из словаря."""
    kind, level = _resolve_block_kind(data)
    blocks: list[ContentBlock] = []

    if kind == "table":
        blocks.append(_build_table_block(data))
    elif kind == "list":
        hint = str(
            _first_present(data, ("type", "kind", "block")) or ""
        ).strip().lower()
        blocks.append(_build_list_block(data, hint))
    elif kind == "caption":
        blocks.append(_build_caption_block(data))
    elif kind == "page_break":
        blocks.append(ContentBlock(kind="page_break"))
    elif kind in SPECIAL_BLOCK_KINDS:
        blocks.append(_build_special_block(data, kind))
    elif kind == "heading":
        text = _as_text(
            _first_present(data, ("text", "content", "value", "title"))
        ).strip()
        if not text and not _explicit_runs(data):
            raise DocxStructureError(f"Блок заголовка (уровень {level}) без текста.")
        if not level:
            level = int(data.get("level") or inherited_level or 1)
        level = max(1, min(9, level))
        blocks.append(
            ContentBlock(
                kind="heading",
                level=level,
                text=text,
                runs=_explicit_runs(data),
                alignment=_as_alignment(_first_present(data, ("alignment", "align"))),
                bold=_as_bool(data.get("bold")),
                italic=_as_bool(data.get("italic")),
            )
        )
    else:  # paragraph
        text = _as_text(_first_present(data, ("text", "content", "value"))).strip()
        if not text and not _explicit_runs(data):
            raise DocxStructureError("Блок paragraph без текста.")
        indent = _first_present(
            data, ("first_line_indent", "first_line_indent_mm", "indent")
        )
        blocks.append(
            ContentBlock(
                kind="paragraph",
                text=text,
                runs=_explicit_runs(data),
                alignment=_as_alignment(_first_present(data, ("alignment", "align"))),
                bold=_as_bool(data.get("bold")),
                italic=_as_bool(data.get("italic")),
                first_line_indent_mm=float(indent) if indent is not None else None,
            )
        )

    # Вложенные дети разворачиваются сразу после самого блока.
    children = _first_present(
        data, ("children", "blocks", "content_blocks", "subblocks", "sub")
    )
    if isinstance(children, (list, tuple)) and children:
        child_level = level + 1 if kind == "heading" else inherited_level
        for child in children:
            blocks.extend(_normalize_item(child, child_level))
    return blocks


def _normalize_item(item: Any, inherited_level: int = 0) -> list[ContentBlock]:
    """Приводит один элемент структуры к списку блоков."""
    if item is None:
        return []
    if isinstance(item, ContentBlock):
        return [item]
    if isinstance(item, dict):
        return _block_from_dict(item, inherited_level)
    if isinstance(item, (list, tuple)):
        result: list[ContentBlock] = []
        for child in item:
            result.extend(_normalize_item(child, inherited_level))
        return result
    if isinstance(item, str):
        text = item.strip()
        if not text:
            return []
        return [ContentBlock(kind="paragraph", text=text)]
    raise DocxStructureError(
        "Элемент структуры должен быть строкой, словарём или списком, "
        f"получено: {type(item).__name__}"
    )


def normalize_structure(structure: Any) -> list[ContentBlock]:
    """Превращает древовидную/списочную структуру в плоский список блоков.

    Поддерживаются:
      * ``list`` блоков;
      * ``dict`` с ключом ``blocks``/``content``/``children`` (и необязательным
        ``title`` — он становится заголовком h1);
      * одиночный блок-словарь, строка, вложенные списки любой глубины.
    """
    if structure is None:
        raise DocxStructureError("Передана пустая структура контента (None).")

    if isinstance(structure, dict):
        blocks: list[ContentBlock] = []
        title = _first_present(structure, ("title", "document_title", "doc_title"))
        if title not in (None, ""):
            blocks.append(ContentBlock(kind="heading", level=1, text=_as_text(title)))
        body = _first_present(
            structure,
            ("blocks", "content", "children", "content_blocks", "sections", "items"),
        )
        known_type = _first_present(structure, ("type", "kind", "block", "block_type"))
        if body is not None and known_type is None:
            for child in _as_list(body):
                blocks.extend(_normalize_item(child, 1))
            if not blocks:
                raise DocxStructureError(
                    "Структура документа пуста: ни title, ни blocks не дали контента."
                )
            return blocks
        # Словарь описывает один блок (возможно, с детьми).
        blocks.extend(_block_from_dict(structure))
        return blocks

    blocks = []
    for child in _as_list(structure):
        blocks.extend(_normalize_item(child))
    if not blocks:
        raise DocxStructureError(
            "Структура документа пуста: передан пустой список или только пустые строки."
        )
    return blocks


def _as_list(value: Any) -> list[Any]:
    """Гарантирует список для итерации (строка не разворачивается в символы)."""
    if isinstance(value, (list, tuple)):
        return list(value)
    return [value]



# --------------------------------------------------------------------------------------
# Низкоуровневые XML-хелперы (python-docx.oxml)
# --------------------------------------------------------------------------------------

#: Порядок детей ``w:tblPr`` по схеме OOXML (для корректной вставки tblBorders).
_TBLPR_TAG_SEQ = (
    "w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual",
    "w:tblStyleRowBandSize", "w:tblStyleColBandSize", "w:tblW", "w:jc",
    "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd", "w:tblLayout",
    "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription",
    "w:tblPrChange",
)

#: Порядок детей ``w:trPr`` по схеме OOXML (для cantSplit / tblHeader).
_TRPR_TAG_SEQ = (
    "w:cnfStyle", "w:divId", "w:gridBefore", "w:gridAfter", "w:wBefore",
    "w:wAfter", "w:cantSplit", "w:trHeight", "w:tblHeader", "w:tblCellSpacing",
    "w:jc", "w:hidden", "w:ins", "w:del", "w:trPrChange",
)

#: Тематические атрибуты шрифта, перебивающие явно заданное имя шрифта.
_RFONT_THEME_ATTRS = ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme")
#: Тематические атрибуты цвета, перебивающие явно заданный RGB.
_COLOR_THEME_ATTRS = ("w:themeColor", "w:themeTint", "w:themeShade")
#: Признаки медиа/полей внутри рана — такие раны нельзя пересобирать по тексту.
_RUN_MEDIA_TAGS = (
    "w:drawing", "w:pict", "w:object", "w:fldChar", "w:instrText",
    "w:footnoteReference", "w:endnoteReference", "w:commentReference",
)


def _insert_in_order(parent: Any, element: Any, tag_seq: Sequence[str]) -> None:
    """Вставляет элемент в ``parent`` согласно порядку тегов схемы OOXML."""
    tag_name = element.tag
    try:
        own_name = next(name for name in tag_seq if qn(name) == tag_name)
    except StopIteration:
        parent.append(element)
        return
    index = list(tag_seq).index(own_name)
    successors = {qn(name) for name in tag_seq[index + 1:]}
    for child in parent:
        if child.tag in successors:
            child.addprevious(element)
            return
    parent.append(element)


def _set_font_on(
    font: Any,
    rPr_owner: Any,
    *,
    font_name: str | None = None,
    size_pt: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    underline: bool | None = None,
) -> None:
    """Универсальная установка шрифта для рана или стиля.

    ``font`` — прокси ``run.font``/``style.font``; ``rPr_owner`` — элемент,
    имеющий ``get_or_add_rPr()`` (``w:r`` у рана, ``w:style`` у стиля).

    Тематические атрибуты (``w:asciiTheme``, ``w:themeColor``) перебивают явные
    значения в стилях ``Heading N`` — без их удаления пресет не применяется.
    """
    if font_name:
        font.name = font_name
        rPr = rPr_owner.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.insert(0, rFonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rFonts.set(qn(attr), font_name)
        for attr in _RFONT_THEME_ATTRS:
            if rFonts.get(qn(attr)) is not None:
                del rFonts.attrib[qn(attr)]

    if size_pt is not None:
        font.size = Pt(size_pt)
        rPr = rPr_owner.get_or_add_rPr()
        sz_cs = rPr.find(qn("w:szCs"))
        if sz_cs is None:
            sz_cs = OxmlElement("w:szCs")
            rPr.append(sz_cs)
        sz_cs.set(qn("w:val"), str(int(round(size_pt * 2))))

    if color:
        rPr = rPr_owner.get_or_add_rPr()
        color_el = rPr.find(qn("w:color"))
        if color == "auto":
            # «auto» = наследовать цвет из стиля: снимаем явное значение.
            if color_el is not None:
                rPr.remove(color_el)
        else:
            font.color.rgb = RGBColor.from_string(color)
            color_el = rPr.find(qn("w:color"))
            if color_el is not None:
                for attr in _COLOR_THEME_ATTRS:
                    if color_el.get(qn(attr)) is not None:
                        del color_el.attrib[qn(attr)]

    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic
    if underline is not None:
        font.underline = underline


def set_run_font(
    run: Run,
    *,
    font_name: str | None = None,
    size_pt: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    underline: bool | None = None,
) -> None:
    """Задаёт шрифт/кегль/цвет/начертание рана (см. :func:`_set_font_on`)."""
    _set_font_on(
        run.font,
        run._r,
        font_name=font_name,
        size_pt=size_pt,
        color=color,
        bold=bold,
        italic=italic,
        underline=underline,
    )


def set_style_font(
    style: Any,
    *,
    font_name: str | None = None,
    size_pt: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    """Задаёт шрифт/кегль/цвет/начертание стиля ``styles.xml``."""
    _set_font_on(
        style.font,
        style.element,
        font_name=font_name,
        size_pt=size_pt,
        color=color,
        bold=bold,
        italic=italic,
    )


def _style_or_none(document: DocumentObject, name: str) -> Any:
    """Возвращает стиль по имени либо ``None``, если его нет в шаблоне."""
    try:
        return document.styles[name]
    except (KeyError, ValueError):
        return None



def apply_line_spacing(pf: Any, line_spacing: float, rule: LineSpacingRule) -> None:
    """Применяет межстрочный интервал с учётом правила.

    Для ``exactly``/``at_least`` значение трактуется как пункты (``Length``),
    иначе — как кратность высоты строки (``float``). Порядок важен: сначала
    задаётся значение, затем правило (иначе python-docx перетрёт значение).
    """
    if rule in POINT_BASED_RULES:
        pf.line_spacing = Pt(line_spacing)
        pf.line_spacing_rule = LINE_RULE_MAP[rule]
    else:
        pf.line_spacing = float(line_spacing)



def apply_paragraph_format(
    pf: Any,
    *,
    alignment: Alignment | None = None,
    space_before: float | None = None,
    space_after: float | None = None,
    line_spacing: float | None = None,
    line_spacing_rule: LineSpacingRule | None = None,
    first_line_indent_mm: float | None = None,
    left_indent_mm: float | None = None,
    right_indent_mm: float | None = None,
    keep_with_next: bool | None = None,
    widow_control: bool | None = None,
    page_break_before: bool | None = None,
) -> None:
    """Накатывает параметры абзаца; ``None`` означает «не трогать»."""
    if alignment is not None:
        pf.alignment = ALIGNMENT_MAP[alignment]
    if space_before is not None:
        pf.space_before = Pt(space_before)
    if space_after is not None:
        pf.space_after = Pt(space_after)
    if line_spacing is not None and line_spacing_rule is not None:
        apply_line_spacing(pf, line_spacing, line_spacing_rule)
    if first_line_indent_mm is not None:
        pf.first_line_indent = Mm(first_line_indent_mm)
    if left_indent_mm is not None:
        pf.left_indent = Mm(left_indent_mm)
    if right_indent_mm is not None:
        pf.right_indent = Mm(right_indent_mm)
    if keep_with_next is not None:
        pf.keep_with_next = keep_with_next
    if widow_control is not None:
        pf.widow_control = widow_control
    if page_break_before is not None:
        pf.page_break_before = page_break_before


def set_table_borders(table: Table, cfg: TableStyle) -> None:
    """Записывает ``w:tblBorders`` (видимые либо ``none``) в правильную позицию."""
    tblPr = table._tbl.tblPr
    existing = tblPr.find(qn("w:tblBorders"))
    if existing is not None:
        tblPr.remove(existing)
    borders = OxmlElement("w:tblBorders")
    value = "single" if cfg.borders else "none"
    size = str(cfg.border_sz_eighths)
    for edge in TABLE_BORDER_EDGES:
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), value)
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), cfg.border_color)
        borders.append(element)
    _insert_in_order(tblPr, borders, _TBLPR_TAG_SEQ)


def _flag_element(tag: str) -> Any:
    """Создаёт флаг-элемент ``w:cantSplit``/``w:tblHeader`` со значением true."""
    element = OxmlElement(tag)
    element.set(qn("w:val"), "true")
    return element


def set_row_properties(
    row: Any, *, cant_split: bool | None = None, tbl_header: bool | None = None
) -> None:
    """Выставляет/снимает ``w:cantSplit`` и ``w:tblHeader`` для строки таблицы."""
    if cant_split is None and tbl_header is None:
        return
    trPr = row._tr.get_or_add_trPr()
    for tag, enabled in (("w:cantSplit", cant_split), ("w:tblHeader", tbl_header)):
        if enabled is None:
            continue
        existing = trPr.find(qn(tag))
        if existing is not None:
            trPr.remove(existing)
        if enabled:
            _insert_in_order(trPr, _flag_element(tag), _TRPR_TAG_SEQ)


def element_has_media(element: Any) -> bool:
    """True, если внутри есть картинки/объекты/поля — их нельзя пересобирать."""
    for tag in _RUN_MEDIA_TAGS:
        qualified = qn(tag)
        if element.find(qualified) is not None:
            return True
        if element.findall(".//" + qualified):
            return True
    return False


def remove_paragraph(paragraph: Paragraph) -> None:
    """Удаляет параграф из тела документа."""
    element = paragraph._p
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def ensure_document_not_empty(document: DocumentObject) -> None:
    """Гарантирует наличие хотя бы одного параграфа (требование OOXML/Word)."""
    body = document.element.body
    has_block = any(
        child.tag in (qn("w:p"), qn("w:tbl")) for child in body.iterchildren()
    )
    if not has_block:
        document.add_paragraph("")



# --------------------------------------------------------------------------------------
# Очистка «грязного» текста и эвристики по пробелам
# --------------------------------------------------------------------------------------


def sanitize_text(text: Any, cleanup: CleanupPolicy) -> str:
    """Вычищает «грязное» оформление текста согласно :class:`CleanupPolicy`.

    Убирает серии пробелов/табуляций (в том числе неразрывные и экзотические
    юникодные пробелы), «хвостовые» пробелы и переводит табы в пробелы.
    Длинные серии подчёркиваний (``____``) сохраняются — это линии для подписей.
    """
    source = _as_text(text)
    if not source:
        return ""
    if cleanup.text_mode != CleanupMode.CLEAN:
        return source

    result = source.replace("\t", " ")
    if cleanup.normalize_nbsp:
        result = result.replace("\u00a0", " ")
    # Экзотические юникодные пробелы -> обычный пробел.
    result = re.sub(r"[\u2000-\u200b\u202f\u205f\u3000]", " ", result)

    limit = max(1, cleanup.max_consecutive_spaces)
    replacement = " " * limit
    # Подчёркивания не трогаем: «____» — линия для ручной подписи.
    protected: list[str] = []

    def _protect(match: re.Match[str]) -> str:
        protected.append(match.group(0))
        return f"\x00{len(protected) - 1}\x00"

    result = _UNDERLINE_GAP_RE.sub(_protect, result)
    result = _MULTISPACE_RE.sub(replacement, result)
    for index, chunk in enumerate(protected):
        result = result.replace(f"\x00{index}\x00", chunk)

    result = re.sub(r" +$", "", result)
    if cleanup.strip_edges:
        result = result.strip()
    return result


def _leading_space_count(text: str) -> int:
    """Число ведущих пробельных символов (пробелы, табы, неразрывные пробелы)."""
    match = _LEADING_WS_RE.match(text.replace("\t", " "))
    return len(match.group(0)) if match else 0


def _line_capacity_chars(preset: StylePreset) -> int:
    """Приблизительная ёмкость строки в символах — нужна для эвристик по пробелам."""
    page = preset.page
    width_mm = A4_LONG_MM if page.is_landscape else A4_SHORT_MM
    text_width = width_mm - page.left - page.right
    if text_width <= 0:
        return _FALLBACK_LINE_CAPACITY
    # 1 pt = 0.3528 мм; средняя ширина символа ~ 0.5 кегля.
    char_mm = max(_AVG_CHAR_WIDTH_RATIO * preset.body.font_size * (25.4 / 72.0), 0.5)
    return max(10, int(text_width / char_mm))


def _detect_shift_alignment(
    raw_text: str, preset: StylePreset, capacity: int
) -> Alignment | None:
    """Определяет «выравнивание пробелами» и возвращает настоящее выравнивание.

    Эвристика (ТЗ: сдвиг текста пробелами заменяется нормальным выравниванием):

    * длинный абзац (занимает > 75 % строки) — ведущие пробелы считаются попыткой
      красной строки: их просто снимают, отступ берётся из пресета (``None``);
    * короткая строка со сдвигом в средней трети — центрирование;
    * короткая строка с большим сдвигом — прижатие вправо.
    """
    cleanup = preset.cleanup
    if not cleanup.fix_space_alignment:
        return None
    leading = _leading_space_count(raw_text)
    if leading < cleanup.alignment_shift_spaces:
        return None

    body = raw_text.strip()
    if not body:
        return None
    # Длинные абзацы не выравниваем по пробелам — это абзацный отступ.
    if len(body) + leading > capacity * 0.75:
        return None

    ratio = leading / max(capacity, 1)
    if 0.15 <= ratio <= 0.45:
        return Alignment.CENTER
    return Alignment.RIGHT


def _effective_columns(header: Sequence[str], rows: Sequence[Sequence[str]]) -> int:
    """Реальная ширина таблицы по шапке и строкам (защита от «рваных» таблиц)."""
    widths = [len(header)] if header else []
    widths.extend(len(row) for row in rows)
    return max(widths) if widths else 0


def _format_caption_text(
    block: ContentBlock, preset: StylePreset
) -> tuple[str, CaptionPlacement]:
    """Собирает текст подписи и её размещение по правилам пресета."""
    captions = preset.captions
    kind = block.caption_kind or "table"
    placement = captions.placement_for(kind)
    body = block.caption or block.text
    return captions.format_text(kind, block.caption_number, body), placement



# --------------------------------------------------------------------------------------
# Распознавание секций документа (титульный лист / реферат / задание / содержание)
# --------------------------------------------------------------------------------------

#: Год в подвале титульного листа: «2026», «2026 г.», «2026 год».
_YEAR_RE = re.compile(r"\b(?:19|20)\d{2}\b")
#: Минимальная длина ключевого слова (короче — слишком общие, пропускаем).
_MIN_KEYWORD_LEN = 3


def _paragraph_breaks_page(paragraph: Paragraph) -> bool:
    """True, если абзац начинает новую страницу или содержит разрыв страницы."""
    element = paragraph._p
    pPr = element.find(qn("w:pPr"))
    if pPr is not None and pPr.find(qn("w:pageBreakBefore")) is not None:
        return True
    for br in element.findall(".//" + qn("w:br")):
        if br.get(qn("w:type")) == "page":
            return True
    return element.find(".//" + qn("w:lastRenderedPageBreak")) is not None


def _normalize_match_text(text: str) -> str:
    """Приводит строку к виду для сравнения с ключевыми словами."""
    cleaned = re.sub(r"[^\w\s]+", " ", text or "", flags=re.UNICODE)
    return re.sub(r"\s+", " ", cleaned).strip().lower()


def _matches_any(text: str, keywords: Sequence[str]) -> bool:
    """True, если начало текста содержит одно из ключевых слов."""
    normalized = _normalize_match_text(text)
    if not normalized:
        return False
    head = normalized[:60]
    return any(
        str(kw).strip().lower() in head
        for kw in keywords
        if str(kw).strip() and len(str(kw).strip()) >= _MIN_KEYWORD_LEN
    )


def _paragraph_has_underline(paragraph: Paragraph) -> bool:
    """True, если в абзаце есть подчёркнутый текст или линия ``____``."""
    if _UNDERLINE_GAP_RE.search(paragraph.text or ""):
        return True
    for run in paragraph.runs:
        if run.font.underline:
            return True
    return False


def _is_year_line(text: str) -> bool:
    """Строка похожа на подвал титульника: содержит год."""
    return bool(_YEAR_RE.search(text or ""))


@dataclass
class DocumentSections:
    """Разбиение абзацев документа на титульный лист и специальные разделы."""

    title_page: list[Paragraph] = field(default_factory=list)
    special: dict[str, list[Paragraph]] = field(default_factory=dict)
    #: Структурные заголовки (ВВЕДЕНИЕ, ЗАКЛЮЧЕНИЕ, СПИСОК ...) — новая страница.
    structural: list[Paragraph] = field(default_factory=list)
    #: Нумерованные подзаголовки (1.1, 1.1.1 ...) — профиль hN БЕЗ новой страницы.
    subheadings: list[tuple[int, Paragraph]] = field(default_factory=list)
    #: индекс абзаца в теле (среди ``w:p``) -> тип секции.
    kind_by_index: dict[int, str] = field(default_factory=dict)

    def kind_for(self, index: int) -> str | None:
        """Тип секции по индексу абзаца (``None`` — обычный текст)."""
        return self.kind_by_index.get(index)

    @property
    def has_special(self) -> bool:
        """Найдены ли специальные разделы."""
        return bool(self.special)


#: Окно поиска титульного листа (в абзацах) — титул не бывает длиннее.
_TITLE_WINDOW = 45
#: Максимальная длина строки подвала титула («Иваново 2026», «Москва, 2026 г.»).
_MAX_FOOTER_LEN = 40


def _has_explicit_page_break(paragraph: Paragraph) -> bool:
    """True, если в абзаце есть ЯВНЫЙ разрыв страницы.

    В отличие от :func:`_paragraph_breaks_page` не учитывает
    ``w:lastRenderedPageBreak`` — это служебная метка Word о том, где страница
    фактически закончилась при последнем рендере. Опора на неё принимала за
    границу титула любой абзац, случайно оказавшийся вверху страницы.
    """
    element = paragraph._p
    pPr = element.find(qn("w:pPr"))
    if pPr is not None:
        pbb = pPr.find(qn("w:pageBreakBefore"))
        # ``w:val="0"/"false"`` — ВЫКЛЮЧЕННЫЙ разрыв (Word оставляет такой тег после
        # сброса свойства), настоящим разрывом он не является.
        if pbb is not None and pbb.get(qn("w:val")) not in ("0", "false", "none"):
            return True
    for br in element.findall(".//" + qn("w:br")):
        if br.get(qn("w:type")) == "page":
            return True
    return False


def _classify_document_sections(
    document: DocumentObject, preset: StylePreset
) -> DocumentSections:
    """Определяет границы титульного листа и спецразделов по структуре документа.

    Титульный лист — абзацы до первого спецраздела или до первого разрыва
    страницы. Спецразделы (реферат/задание/содержание) начинаются со своего
    заголовка-ключа и тянутся до начала следующего спецраздела.
    """
    result = DocumentSections()
    special_cfg = preset.special_sections
    title_cfg = preset.title_page

    paragraphs = [
        Paragraph(child, document)
        for child in document.element.body.iterchildren()
        if child.tag == qn("w:p")
    ]
    if not paragraphs:
        return result

    # 1. Первый спецраздел — только по СТРОГОМУ совпадению названия раздела.
    #    Прежний поиск подстроки принимал строку подписи «Задание приняла к
    #    исполнению … Петрова П.П.» за начало раздела «ЗАДАНИЕ», из-за чего
    #    титульный лист обрезался на подписях, а «город + год» уезжал на лист 2.
    first_special: int | None = None
    if special_cfg.enabled:
        for index, paragraph in enumerate(paragraphs):
            for kind in SPECIAL_SECTION_KINDS:
                if kind not in special_cfg.apply_to:
                    continue
                if _is_section_title_line(paragraph.text, special_cfg.titles_for(kind)):
                    first_special = index
                    break
            if first_special is not None:
                break

    # 2. Явный разрыв страницы и блок подвала («город + год») внутри титула.
    scan_end = first_special if first_special is not None else len(paragraphs)
    page_break_at: int | None = None
    for index in range(scan_end):
        if _has_explicit_page_break(paragraphs[index]):
            page_break_at = index
            break

    search_end = (
        first_special
        if first_special is not None
        else min(len(paragraphs), _TITLE_WINDOW)
    )
    footer_at: int | None = None
    for index in range(search_end):
        text = paragraphs[index].text.strip()
        if text and len(text) <= _MAX_FOOTER_LEN and _is_year_line(text):
            footer_at = index

    # Граница титула — ближайшая из достоверных. Разрыв, стоящий ПЕРЕД подвалом,
    # границей не считается: титул продлевается до конца блока «город + год», а
    # сам разрыв снимается при вёрстке (разрыв должен идти ПОСЛЕ года).
    candidates: list[int] = []
    if first_special is not None:
        candidates.append(first_special)
    if page_break_at is not None and (footer_at is None or page_break_at > footer_at):
        candidates.append(page_break_at + 1)
    if footer_at is not None:
        candidates.append(footer_at + 1)
    cut = min(candidates) if candidates else 0

    # 3. Титульный лист применяем только при явных признаках: разрыв страницы,
    #    следующий спецраздел, подвал с годом или подписи на первой странице.
    title_evidence = page_break_at is not None or first_special is not None
    if not title_evidence and footer_at is not None:
        title_evidence = True
    if not title_evidence and title_cfg.enabled:
        for paragraph in paragraphs[:cut]:
            text = paragraph.text.strip()
            if not text:
                continue
            if _is_year_line(text) or _paragraph_has_underline(paragraph):
                title_evidence = True
                break
            if _matches_any(text, title_cfg.signature_keywords):
                title_evidence = True
                break

    if title_cfg.enabled and title_evidence and cut > 0:
        for index in range(cut):
            result.title_page.append(paragraphs[index])
            result.kind_by_index[index] = "title_page"

    # 4. Спецразделы и заголовки разделов. Структурный заголовок (ВВЕДЕНИЕ,
    #    ЗАКЛЮЧЕНИЕ, СПИСКИ, ПРИЛОЖЕНИЯ) закрывает спецраздел и начинается с
    #    новой страницы. Нумерованные заголовки распознаются только ВНЕ
    #    спецразделов, чтобы пункт «1. Содержание работы:» бланка задания не
    #    превращался в заголовок раздела.
    current: str | None = None
    #: Шапка первого бланка (ведомство/вуз ПЕРЕД заголовком «ЗАДАНИЕ») лежит между
    #: концом титульного листа и первым спецразделом. Она относится к этому разделу,
    #: иначе остаётся «сиротой»: получает стиль основного текста и теряет разрыв
    #: страницы, а заголовок «ЗАДАНИЕ» уезжает на следующую страницу.
    leading: list[tuple[int, Paragraph]] = []
    first_special_seen = False
    for index in range(cut, len(paragraphs)):
        paragraph = paragraphs[index]
        text = paragraph.text.strip()
        matched: str | None = None
        if special_cfg.enabled:
            for kind in SPECIAL_SECTION_KINDS:
                if kind not in special_cfg.apply_to:
                    continue
                if _is_section_title_line(text, special_cfg.titles_for(kind)):
                    matched = kind
                    break
        if matched is not None:
            current = matched
            target = result.special.setdefault(current, [])
            if not first_special_seen:
                first_special_seen = True
                # Шапка бланка задания (ведомство/вуз) идёт ПЕРЕД заголовком
                # «ЗАДАНИЕ»; относим её к этому разделу, а не к «сиротам».
                # У реферата/содержания шапки нет — для них leading не трогаем.
                if matched == "task":
                    for _index, _paragraph in leading:
                        target.append(_paragraph)
                        result.kind_by_index[_index] = current
            leading = []
            target.append(paragraph)
            result.kind_by_index[index] = current
            continue
        if text:
            level = _section_heading_level(
                text, preset, allow_numbered=current is None
            )
            if level == 1:
                current = None
                leading = []
                result.structural.append(paragraph)
                result.kind_by_index[index] = "structure"
                continue
            if level is not None and current is None:
                leading = []
                result.subheadings.append((level, paragraph))
                result.kind_by_index[index] = f"heading{level}"
                continue
        if current is None:
            # Возможная шапка первого бланка (до встречи его заголовка). Пустые
            # абзацы (в т.ч. разрывы страницы) не трогаем — их разберёт чистка
            # распорок в _sanitize_body.
            if text:
                leading.append((index, paragraph))
            continue
        # Пустой абзац-разрыв страницы закрывает текущий спецраздел.
        starts_new_page = _paragraph_breaks_page(paragraph) and not text
        if starts_new_page and result.special.get(current):
            current = None
            continue
        result.special.setdefault(current, []).append(paragraph)
        result.kind_by_index[index] = current
    return result


def _set_runs_case(paragraph: Paragraph, case: TextCase) -> None:
    """Меняет регистр текста абзаца (для заголовков спецразделов)."""
    if case == TextCase.NORMAL:
        return
    for run in paragraph.runs:
        if not run.text or element_has_media(run._r):
            continue
        if case == TextCase.UPPERCASE:
            run.text = run.text.upper()


def _apply_text_block_style(
    paragraph: Paragraph,
    *,
    font_name: str,
    font_size: float,
    color: str,
    bold: bool | None,
    italic: bool | None,
    alignment: Alignment,
    space_before: float | None = None,
    space_after: float | None = None,
    first_line_indent_mm: float = 0.0,
    left_indent_mm: float | None = None,
    line_spacing: float | None = None,
    line_spacing_rule: LineSpacingRule | None = None,
    keep_with_next: bool | None = None,
    page_break_before: bool | None = None,
) -> None:
    """Накатывает оформление на абзац, НЕ меняя его текст и раны.

    Используется для титульного листа и спецразделов: сохраняет дефисы, буквы
    и подчёркивания исходного текста, но приводит шрифт/интервалы/отступы к пресету.
    """
    for run in paragraph.runs:
        if element_has_media(run._r):
            continue
        set_run_font(
            run,
            font_name=font_name,
            size_pt=font_size,
            color=color,
            bold=bold,
            italic=italic,
        )
    apply_paragraph_format(
        paragraph.paragraph_format,
        alignment=alignment,
        space_before=space_before,
        space_after=space_after,
        line_spacing=line_spacing,
        line_spacing_rule=line_spacing_rule,
        first_line_indent_mm=first_line_indent_mm,
        left_indent_mm=left_indent_mm,
        keep_with_next=keep_with_next,
        page_break_before=page_break_before,
    )


def _footer_start_index(
    paragraphs: list[Paragraph],
    signature_keywords: Sequence[str] = (),
) -> int | None:
    """Индекс первого абзаца подвала титульного листа («Город, Год»).

    Подвал — это только блок «город + год» в самом конце титула. Раньше функция
    жадно включала в подвал ещё и предыдущий короткий абзац, из-за чего строка
    подписи («Петрова П.П. ______________») получала вертикальную распорку
    подвала (240 pt) ВМЕСТЕ с «Иваново, 2026 г.». Две распорки по 240 pt
    (~169 мм) не помещались в ~257 мм текста A4, титульный лист перетекал на
    вторую страницу, а «ЗАДАНИЕ» уезжало на третью.
    """
    keywords = tuple(str(k).strip().lower() for k in signature_keywords if str(k).strip())
    nonempty = [i for i, p in enumerate(paragraphs) if p.text.strip()]
    if not nonempty:
        return None
    last = nonempty[-1]
    if not _is_year_line(paragraphs[last].text):
        return None
    start = last
    if len(nonempty) >= 2:
        previous = nonempty[-2]
        previous_paragraph = paragraphs[previous]
        previous_text = previous_paragraph.text.strip()
        # Строку подписи в подвал не тянем: она относится к блоку подписей, а не
        # к «город + год», и распорка подвала ей не положена.
        is_signature = _paragraph_has_underline(previous_paragraph) or _matches_any(
            previous_text, keywords
        )
        if (
            len(previous_text) <= 40
            and not _is_year_line(previous_text)
            and not is_signature
        ):
            start = previous
    return start


def _strip_auto_numbering(paragraph: Paragraph) -> bool:
    """Снимает с абзаца автоматическую нумерацию/маркер списка (``w:numPr``).

    Автозамена Word превращает строки вида «АО — акционерное общество» и
    «Цель работы — анализ…» в элементы нумерованного списка. По ГОСТ 7.32 такие
    строки должны оставаться обычными абзацами, поэтому привязку к списку
    удаляем, а сам текст не трогаем.
    """
    pPr = paragraph._p.find(qn("w:pPr"))
    if pPr is None:
        return False
    numPr = pPr.find(qn("w:numPr"))
    if numPr is None:
        return False
    pPr.remove(numPr)
    return True


def _set_right_tab_stop(paragraph: Paragraph, pos_mm: float, leader: str) -> None:
    """Ставит правую позицию табуляции с заполнителем (точки/черта) у правого поля."""
    pf = paragraph.paragraph_format
    pf.tab_stops.clear_all()
    pf.tab_stops.add_tab_stop(
        Mm(pos_mm),
        WD_TAB_ALIGNMENT.RIGHT,
        _TAB_LEADER_MAP.get(str(leader or "none").lower(), WD_TAB_LEADER.SPACES),
    )


def _paragraph_has_tab(paragraph: Paragraph) -> bool:
    """True, если в абзаце есть СИМВОЛ табуляции.

    Важно искать именно ``w:r/w:tab`` (символ внутри рана), а не ``w:tabs``:
    ``w:pPr/w:tabs/w:tab`` — это ОПРЕДЕЛЕНИЕ позиции табуляции, которое
    добавляет сам движок через :func:`_set_right_tab_stop`. Прежний поиск
    ``.//w:tab`` находил определения позиций и принимал любой абзац с правой
    табуляцией (например, пункт списка в бланке задания) за строку подписи,
    из-за чего ему ставили вторую позицию табуляции и выравнивание влево.
    """
    for run_element in paragraph._p.findall(qn("w:r")):
        if run_element.find(qn("w:tab")) is not None:
            return True
    return False


def _structural_titles(preset: StylePreset) -> list[str]:
    """Ключевые слова структурных разделов из пресета (или значения по умолчанию)."""
    titles = getattr(preset.special_sections, "structural_titles", None)
    return list(titles) if titles else list(_DEFAULT_STRUCTURAL_TITLES)


def _is_structural_title(text: str, preset: StylePreset) -> bool:
    """True, если строка — заголовок структурного раздела (ВВЕДЕНИЕ, СПИСОК и т.п.)."""
    return _matches_any(text, _structural_titles(preset))


def _looks_like_toc_entry(text: str) -> bool:
    """True, если строка похожа на строку содержания «Заголовок …… 5»."""
    stripped = (text or "").strip()
    if not stripped:
        return False
    return bool(re.match(r"^.*?[\s.\u2026]+\s*\d{1,4}\s*$", stripped))


def _upper_ratio(text: str) -> float:
    """Доля заглавных букв среди всех букв строки."""
    letters = [ch for ch in text if ch.isalpha()]
    if not letters:
        return 0.0
    return sum(1 for ch in letters if ch.isupper()) / len(letters)


def _heading_shape_ok(stripped: str) -> bool:
    """Отсекает предложения и строки содержания, оставляя «форму» заголовка."""
    if not stripped or len(stripped) > 100:
        return False
    if _looks_like_toc_entry(stripped):
        return False
    if _upper_ratio(stripped) >= 0.6:
        return True
    # Не-заглавные строки: короткие и без завершающей «предложенческой» точки.
    # Двоеточие в конце тоже отметаем: «1. Содержание работы:» — это пункт
    # бланка задания, а не заголовок раздела с новой страницы.
    if stripped.endswith((".", "!", "?", "\u2026", ":")):
        return False
    return len(stripped.split()) <= _MAX_HEADING_WORDS


def _section_heading_level(
    text: str, preset: StylePreset, *, allow_numbered: bool = True
) -> int | None:
    """Уровень заголовка раздела (1..9) или ``None``.

    Распознаёт два вида заголовков:

    * структурный — «ВВЕДЕНИЕ», «РЕФЕРАТ», «СОДЕРЖАНИЕ», «ЗАКЛЮЧЕНИЕ»,
      «СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ» (по ключевым словам пресета);
    * нумерованный — «1 ГЛАВА. АНАЛИЗ…», «1. АНАЛИЗ…», «1.1 Описание…».
      Глубина номера задаёт уровень: «1» -> h1, «1.1» -> h2, «1.1.1» -> h3.

    Обычные предложения, строки содержания и длинные абзацы заголовками не
    считаются. ``allow_numbered=False`` оставляют только ключевые слова —
    это безопасно внутри специальных разделов, где «1. Содержание работы:»
    является пунктом бланка задания, а не заголовком раздела.
    """
    stripped = (text or "").strip()
    if not _heading_shape_ok(stripped):
        return None
    if _is_structural_title(stripped, preset):
        return 1
    if not allow_numbered:
        return None
    match = _SECTION_NUMBER_RE.match(stripped)
    if not match:
        return None
    parts = [part for part in match.group("number").split(".") if part]
    return max(1, min(len(parts), 9))


def _is_section_title_line(text: str, keywords: Sequence[str]) -> bool:
    """True, если строка — НАЗВАНИЕ специального раздела («РЕФЕРАТ», «ЗАДАНИЕ»).

    :func:`_matches_any` ищет ключевое слово где угодно в первых 60 символах,
    поэтому строка подписи «Задание приняла к исполнению … Петрова П.П.»
    распознавалась как начало раздела «ЗАДАНИЕ». Из-за этого граница титульного
    листа обрезалась на строке подписи, а блок «город + год» выпадал из титула
    и уезжал на вторую страницу.

    Здесь требование жёстче: строка должна быть похожа на заголовок (короткая,
    без завершающей точки) и либо совпадать с ключевым словом целиком, либо
    начинаться с него.
    """
    stripped = (text or "").strip()
    if not _heading_shape_ok(stripped):
        return False
    normalized = _normalize_match_text(stripped)
    if not normalized:
        return False
    for keyword in keywords:
        word = str(keyword or "").strip().lower()
        if not word or len(word) < _MIN_KEYWORD_LEN:
            continue
        if normalized == word or normalized.startswith(word + " "):
            return True
    return False


def _has_explicit_list_marker(text: str) -> bool:
    """True, если в НАЧАЛЕ строки исходника явно стоит маркер списка.

    Только явный маркер («1.», «2)», «а)», «-», «–», «•») даёт право считать
    абзац элементом списка. Если маркера нет, автоматическую нумерацию Word
    (``w:numPr``) снимаем: иначе в реферате после двоеточия появлялись цифры,
    которых в исходном документе не было.
    """
    return bool(_LIST_MARKER_RE.match((text or "").strip()))


def _split_list_marker(text: str) -> tuple[str, str] | None:
    """Разбирает строку на «маркер» и «текст без маркера».

    Возвращает ``None``, если маркера нет. Используется перед наложением маркера
    из конфигурации, чтобы исключить сдвоенные символы («- - пункт»).
    """
    match = _LEADING_MARKER_RE.match((text or "").strip())
    if not match:
        return None
    marker = match.group("marker").strip()
    rest = (match.group("rest") or "").strip()
    if not marker or not rest:
        return None
    return marker, rest


def _clear_page_breaks(paragraph: Paragraph) -> bool:
    """Убирает из абзаца явные разрывы страницы.

    Титульный лист должен занимать ровно первый печатный лист, а разрыв страницы
    ставится строго ПОСЛЕ блока года. Авторы часто ставят ручной разрыв ПЕРЕД
    строкой «Иваново 2026», чтобы «красиво» её перенести, — такой разрыв
    выталкивал город и год на отдельный лист.
    """
    element = paragraph._p
    changed = False
    pPr = element.find(qn("w:pPr"))
    if pPr is not None:
        for flag in pPr.findall(qn("w:pageBreakBefore")):
            pPr.remove(flag)
            changed = True
    for br in element.findall(".//" + qn("w:br")):
        if br.get(qn("w:type")) == "page":
            parent = br.getparent()
            if parent is not None:
                parent.remove(br)
                changed = True
    return changed


def _format_signature_paragraph(
    paragraph: Paragraph,
    cfg: Any,
    space_before: float | None = None,
    space_after: float | None = None,
) -> None:
    """Форматирует строку подписи: сохраняет табуляцию, черта — у правого поля.

    Если в строке есть табуляция, текст до неё остаётся слева (выравнивание
    влево), а часть после табуляции прижимается к правому полю правой позицией
    табуляции с заполнителем. Если табуляции нет — блок смещается вправо отступом.

    ``space_before``/``space_after`` переопределяют интервалы блока: титульный
    лист задаёт их только для первого/последнего абзаца блока, чтобы отступ не
    умножался на каждую строку и титул не переполнял страницу A4.
    """
    if _paragraph_has_tab(paragraph) and getattr(cfg, "preserve_tabs", True):
        _set_right_tab_stop(
            paragraph, getattr(cfg, "tab_stop_mm", 165.0), getattr(cfg, "tab_leader", "underscore")
        )
        alignment = Alignment.LEFT
        left_indent = None
    else:
        alignment = cfg.alignment
        left_indent = cfg.left_indent
    _apply_text_block_style(
        paragraph,
        font_name=cfg.font_name,
        font_size=cfg.font_size,
        color=cfg.color,
        bold=cfg.bold or None,
        italic=cfg.italic or None,
        alignment=alignment,
        space_before=cfg.space_before if space_before is None else space_before,
        space_after=cfg.space_after if space_after is None else space_after,
        first_line_indent_mm=0.0,
        left_indent_mm=left_indent,
        keep_with_next=False,
    )


def _rebuild_toc_entry_text(paragraph: Paragraph) -> str | None:
    """Пересобирает строку содержания в вид «Заголовок<таб>номер».

    Возвращает новый текст (или ``None``, если номер страницы не распознан).
    """
    text = (paragraph.text or "").rstrip()
    match = re.match(r"^(.*?)[\s.\u2026]+\s*(\d{1,4})\s*$", text)
    if match is None:
        return None
    title = match.group(1).rstrip(" .\t\u2026")
    page = match.group(2)
    if not title:
        return None
    new_text = f"{title}\t{page}"
    if new_text == paragraph.text:
        return new_text
    for run in list(paragraph.runs):
        run._r.getparent().remove(run._r)
    paragraph.add_run(new_text)
    return new_text


def _format_toc_entry(paragraph: Paragraph, preset: StylePreset, stats: dict[str, Any]) -> None:
    """Оформляет строку содержания: правая табуляция с точками у правого поля."""
    cfg = preset.special_sections
    if cfg.toc_rebuild_entries:
        _rebuild_toc_entry_text(paragraph)
    _set_right_tab_stop(paragraph, cfg.toc_right_tab_mm, cfg.toc_leader)
    _apply_text_block_style(
        paragraph,
        font_name=cfg.body_font_name,
        font_size=cfg.body_font_size,
        color=preset.body.color,
        bold=None,
        italic=None,
        alignment=cfg.toc_entry_alignment,
        space_before=0.0,
        space_after=0.0,
        first_line_indent_mm=0.0,
        left_indent_mm=0.0,
        line_spacing=cfg.body_line_spacing,
        line_spacing_rule=cfg.body_line_spacing_rule,
    )
    stats["toc_entries"] = stats.get("toc_entries", 0) + 1


def _format_structure_heading(paragraph: Paragraph, preset: StylePreset, stats: dict[str, Any]) -> None:
    """Оформляет структурный заголовок раздела профилем ПЕРВОГО уровня.

    Кегль, начертание и выравнивание берутся из ``preset.heading(1)`` — то есть
    из профиля, который модуль экстракции извлёк из исходного документа. Раньше
    здесь использовались ``special_sections.title_*``, и структурные заголовки
    («ВВЕДЕНИЕ», «ЗАКЛЮЧЕНИЕ», «СПИСОК …») рендерились кеглем основного текста,
    если пресет не подхватывал размер из спецраздела.

    Из ``special_sections`` остаются только интервалы, регистр и разрыв страницы:
    это параметры именно структурных разделов, а не шрифтовой профиль.
    """
    cfg = preset.special_sections
    h1 = preset.heading(1)
    _apply_text_block_style(
        paragraph,
        font_name=h1.font_name,
        font_size=h1.font_size,
        color=h1.color,
        bold=h1.bold,
        italic=h1.italic,
        alignment=h1.alignment,
        space_before=cfg.title_space_before,
        space_after=cfg.title_space_after,
        first_line_indent_mm=h1.first_line_indent,
        line_spacing=cfg.body_line_spacing,
        line_spacing_rule=cfg.body_line_spacing_rule,
        keep_with_next=True,
        page_break_before=cfg.page_break_before,
    )
    _set_runs_case(paragraph, h1.text_case)
    stats["structural_headings"] = stats.get("structural_headings", 0) + 1


def _format_subheading(
    paragraph: Paragraph, level: int, preset: StylePreset, stats: dict[str, Any]
) -> None:
    """Оформляет нумерованный подзаголовок (1.1, 1.1.1 …) профилем ``hN``.

    В отличие от структурного заголовка НЕ ставит разрыв страницы: подразделы
    идут подряд внутри главы.
    """
    cfg = preset.heading(max(1, min(level, 9)))
    _apply_text_block_style(
        paragraph,
        font_name=cfg.font_name,
        font_size=cfg.font_size,
        color=cfg.color,
        bold=cfg.bold,
        italic=cfg.italic,
        alignment=cfg.alignment,
        space_before=cfg.space_before,
        space_after=cfg.space_after,
        first_line_indent_mm=cfg.first_line_indent,
        line_spacing=preset.body.line_spacing,
        line_spacing_rule=preset.body.line_spacing_rule,
        keep_with_next=cfg.keep_with_next,
        page_break_before=False,
    )
    _set_runs_case(paragraph, cfg.text_case)
    stats["subheadings"] = stats.get("subheadings", 0) + 1


#: Пунктов в миллиметре (для перевода кегля в высоту строки).
_PT_PER_MM = 72.0 / 25.4
#: Множитель высоты строки титула относительно кегля (интервал ≈ одинарный).
_TITLE_LINE_FACTOR = 1.2


def _title_text_height_mm(
    paragraphs: list[Paragraph], sizes: list[float], preset: StylePreset
) -> float:
    """Оценка высоты самого текста титула в миллиметрах (без интервалов)."""
    page = preset.page
    long_side = A4_LONG_MM if not page.is_landscape else A4_SHORT_MM
    width_mm = max(
        20.0,
        (long_side if not page.is_landscape else A4_SHORT_MM)
        - page.left
        - page.right,
    )
    total = 0.0
    for paragraph, size in zip(paragraphs, sizes):
        text = (paragraph.text or "").strip()
        font_size = size if size and size > 0 else preset.body.font_size
        char_mm = max(0.1, font_size * _AVG_CHAR_WIDTH_RATIO / _PT_PER_MM)
        per_line = max(8, int(width_mm / char_mm))
        lines = max(1, -(-len(text) // per_line)) if text else 1
        total += lines * font_size * _TITLE_LINE_FACTOR / _PT_PER_MM
    return total


def _title_spacing_scale(
    paragraphs: list[Paragraph],
    sizes: list[float],
    before: list[float],
    after: list[float],
    preset: StylePreset,
    document: DocumentObject | None,
) -> float:
    """Коэффициент сжатия вертикальных отступов титула (0.0..1.0).

    Титульный лист обязан помещаться ровно на первом печатном листе. Считаем
    доступную высоту текста (высота листа минус верхнее и нижнее поля, извлечённые
    из документа) и сравниваем её с суммой «высота текста + все интервалы». Если
    сумма больше, интервалы пропорционально урезаются — так избыточные
    ``space_before``/``space_after`` блоков титульника и паразитные пустые строки
    больше не выталкивают город и год на вторую страницу.
    """
    spacing_pt = sum(before) + sum(after)
    if spacing_pt <= 0:
        return 1.0

    page = preset.page
    height_mm = A4_LONG_MM if not page.is_landscape else A4_SHORT_MM
    if document is not None and document.sections:
        section = document.sections[0]
        try:
            if section.page_height is not None:
                height_mm = section.page_height.mm
        except (AttributeError, ValueError):  # noqa: PERF203
            pass
    available_mm = max(20.0, height_mm - page.top - page.bottom)

    text_mm = _title_text_height_mm(paragraphs, sizes, preset)
    spacing_mm = spacing_pt / _PT_PER_MM
    if text_mm + spacing_mm <= available_mm:
        return 1.0
    return max(0.0, (available_mm - text_mm) / spacing_mm)


def _format_title_page(
    paragraphs: list[Paragraph],
    preset: StylePreset,
    stats: dict[str, Any],
    document: DocumentObject | None = None,
) -> int:
    """Форматирует титульный лист: шапка / заголовок / подписи / подвал.

    К титульному листу НЕ применяется ``body.first_line_indent``; блоки
    распределяются по вертикали за счёт интервалов (подвал прижимается к низу
    значением ``title_page.footer.space_before``).

    Вертикальную распорку (``space_before``) получает ТОЛЬКО первый абзац блока,
    интервал «после блока» — только последний. Суммарные вертикальные отступы
    дополнительно ограничиваются доступной высотой первого листа
    (:func:`_title_spacing_scale`), поэтому титул гарантированно помещается ровно
    на одном печатном листе A4 с учётом полей, извлечённых из документа.

    Явные разрывы страницы внутри титула снимаются: разрыв должен идти строго
    ПОСЛЕ блока года, а не перед ним.
    """
    cfg = preset.title_page
    if not cfg.enabled or not paragraphs:
        return 0

    footer_start = _footer_start_index(paragraphs, cfg.signature_keywords)

    # 1. Размечаем блок каждого абзаца — это позволяет отличить «первый абзац
    #    блока» от последующих и не умножать вертикальную распорку.
    kinds: list[str] = []
    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        if not text:
            kinds.append("")
        elif footer_start is not None and index >= footer_start:
            kinds.append("footer")
        elif _matches_any(text, cfg.signature_keywords) or _paragraph_has_underline(
            paragraph
        ):
            kinds.append("signatures")
        elif _matches_any(text, cfg.title_keywords):
            kinds.append("title")
        else:
            kinds.append("header")

    blocks: dict[str, Any] = {
        "header": cfg.header,
        "title": cfg.title,
        "signatures": cfg.signatures,
        "footer": cfg.footer,
    }

    # 2. План интервалов: распорка — первому абзацу блока, «после» — последнему.
    sizes: list[float] = []
    before: list[float] = []
    after: list[float] = []
    for index, paragraph in enumerate(paragraphs):
        kind = kinds[index]
        block = blocks.get(kind)
        sizes.append(float(getattr(block, "font_size", 0.0) or 0.0))
        if not kind:
            before.append(0.0)
            after.append(0.0)
            continue
        first_in_block = index == 0 or kinds[index - 1] != kind
        last_in_block = index == len(kinds) - 1 or kinds[index + 1] != kind
        want_before = float(getattr(block, "space_before", 0.0) or 0.0)
        if kind == "footer" and not cfg.distribute_vertically:
            want_before = 0.0
        before.append(want_before if first_in_block else 0.0)
        after.append(
            float(getattr(block, "space_after", 0.0) or 0.0)
            if last_in_block
            else 0.0
        )

    # 3. Строго ограничиваем суммарные отступы высотой первого листа.
    scale = _title_spacing_scale(paragraphs, sizes, before, after, preset, document)
    if scale < 1.0:
        stats["title_page_spacing_scaled"] = round(scale, 3)

    touched = 0
    for index, paragraph in enumerate(paragraphs):
        kind = kinds[index]
        # Разрыв страницы внутри титула снимаем: он должен стоять ПОСЛЕ года.
        if _clear_page_breaks(paragraph):
            stats["title_page_breaks_removed"] = (
                stats.get("title_page_breaks_removed", 0) + 1
            )
        if not kind:
            continue
        block = blocks[kind]
        space_before = before[index] * scale
        space_after = after[index] * scale
        if kind == "signatures":
            _format_signature_paragraph(
                paragraph,
                cfg.signatures,
                space_before=space_before,
                space_after=space_after,
            )
        else:
            _apply_text_block_style(
                paragraph,
                font_name=block.font_name,
                font_size=block.font_size,
                color=block.color,
                bold=block.bold or None,
                italic=block.italic or None,
                alignment=block.alignment,
                space_before=space_before,
                space_after=space_after,
                first_line_indent_mm=0.0,
                keep_with_next=False if kind == "footer" else None,
            )
        touched += 1
    stats["title_page_formatted"] = stats.get("title_page_formatted", 0) + touched
    return touched


def _merge_task_signature_pair(
    paragraphs: list[Paragraph], preset: StylePreset, stats: dict[str, Any]
) -> list[Paragraph]:
    """Склеивает пару подписей бланка задания в одну строку из двух колонок.

    «Задание приняла к исполнению ______» и «Руководитель ВКР ______» должны
    стоять ровно: левая часть у левого поля, правая — у правого. Отдельными
    абзацами с отступом 85 мм они выглядели «лесенкой», поэтому два соседних
    абзаца подписей объединяются в один через символ табуляции, а правая позиция
    табуляции ставится у правого поля (~165 мм).

    Возвращает список получившихся слитных абзацев — им позже повторно ставят
    позицию табуляции без заполнителя (это две колонки, а не линия подписи).
    """
    cfg = preset.title_page
    keywords = cfg.signature_keywords
    tab_mm = float(getattr(cfg.signatures, "tab_stop_mm", 165.0) or 165.0)

    def _is_signature(paragraph: Paragraph, text: str) -> bool:
        return _matches_any(text, keywords) or _paragraph_has_underline(paragraph)

    merged: list[Paragraph] = []
    index = 0
    while index < len(paragraphs) - 1:
        left = paragraphs[index]
        right = paragraphs[index + 1]
        left_text = left.text.strip()
        right_text = right.text.strip()
        if not left_text or not right_text:
            index += 1
            continue
        if not (_is_signature(left, left_text) and _is_signature(right, right_text)):
            index += 1
            continue
        if _paragraph_has_tab(left) or _paragraph_has_tab(right):
            # Табуляция уже есть — строки сами встанут по позициям, не трогаем.
            index += 2
            continue
        for run in list(right.runs):
            run._r.getparent().remove(run._r)
        left.add_run("\t")
        left.add_run(right_text)
        # Заполнитель не нужен: это две колонки, а не линия до номера страницы.
        _set_right_tab_stop(left, tab_mm, "none")
        remove_paragraph(right)
        paragraphs.pop(index + 1)
        merged.append(left)
        index += 1

    if merged:
        stats["task_signatures_merged"] = (
            stats.get("task_signatures_merged", 0) + len(merged)
        )
    return merged


def _format_special_section(
    kind: str, paragraphs: list[Paragraph], preset: StylePreset, stats: dict[str, Any]
) -> int:
    """Форматирует специальный раздел: название строго по центру, текст по ширине.

    Названия («РЕФЕРАТ», «ЗАДАНИЕ») — по центру, без красной строки. Шапка бланка
    задания (строки ведомства/вуза ПЕРЕД заголовком) — тоже по центру, но компактно,
    одиночным интервалом и без абзацного отступа. Основной текст — ``justify`` +
    абзацный отступ. Пункты списка сохраняют дефисы/буквы и получают отступ списка.

    Разрыв страницы ставится ТОЛЬКО на первом абзаце раздела: для задания это шапка
    бланка, для реферата/содержания — сам заголовок. Раньше разрыв вешался на
    заголовок всегда, поэтому при наличии шапки бланка «ЗАДАНИЕ» уезжало на
    следующую страницу, а шапка оставалась на предыдущей.
    """
    cfg = preset.special_sections
    if not cfg.enabled or kind not in cfg.apply_to or not paragraphs:
        return 0

    # Подписи бланка задания («Задание приняла…» / «Руководитель…») выравниваем
    # в две колонки одной строкой через правую табуляцию.
    merged_signatures: list[Paragraph] = []
    if kind == "task":
        merged_signatures = _merge_task_signature_pair(paragraphs, preset, stats)

    # Бланк задания — компактная форма: в оригинале межстрочный интервал одинарный,
    # а межабзацные отступы сведены к нулю. Полуторный интервал основного текста
    # здесь раздувал бы бланк и выталкивал подписи на следующую страницу.
    is_blank = kind == "task"
    line_spacing = 1.0 if is_blank else cfg.body_line_spacing
    line_spacing_rule = LineSpacingRule.SINGLE if is_blank else cfg.body_line_spacing_rule

    # Заголовок раздела определяем по ключевому слову, а не по «первому абзацу»:
    # у бланка задания перед «ЗАДАНИЕ» идёт шапка ведомства/вуза.
    titles = cfg.titles_for(kind)
    first_index: int | None = None
    title_index: int | None = None
    for index, paragraph in enumerate(paragraphs):
        if not paragraph.text.strip():
            continue
        if first_index is None:
            first_index = index
        if title_index is None and _is_section_title_line(paragraph.text, titles):
            title_index = index
    if first_index is None:
        return 0
    if title_index is None:
        title_index = first_index

    # Снимаем явные разрывы страницы в шапке бланка (и в пустых абзацах-распорках
    # перед ней): разрыв ставится ровно один — на первом абзаце раздела. Иначе
    # пустой абзац с ``<w:br type="page">`` плюс ``page_break_before`` шапки давали
    # бы лишнюю пустую страницу.
    if first_index < title_index:
        for paragraph in paragraphs[:title_index]:
            _clear_page_breaks(paragraph)

    touched = 0
    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue

        # Разрыв страницы — только на ПЕРВОМ абзаце раздела (шапке бланка или
        # заголовке), а не на самом заголовке при наличии шапки.
        starts_page = cfg.page_break_before and index == first_index

        # Шапка бланка задания (всё до заголовка) — по центру, компактно.
        if index < title_index:
            _apply_text_block_style(
                paragraph,
                font_name=cfg.title_font_name,
                font_size=cfg.title_font_size,
                color=cfg.title_color,
                bold=False,
                italic=False,
                alignment=Alignment.CENTER,
                space_before=0.0,
                space_after=0.0,
                first_line_indent_mm=0.0,
                line_spacing=line_spacing,
                line_spacing_rule=line_spacing_rule,
                keep_with_next=True,
                page_break_before=starts_page,
            )
            touched += 1
            continue

        # Заголовок раздела.
        if index == title_index:
            _apply_text_block_style(
                paragraph,
                font_name=cfg.title_font_name,
                font_size=cfg.title_font_size,
                color=cfg.title_color,
                bold=cfg.title_bold,
                italic=cfg.title_italic,
                alignment=cfg.title_alignment,
                space_before=cfg.title_space_before,
                space_after=cfg.title_space_after,
                first_line_indent_mm=cfg.title_first_line_indent,
                line_spacing=line_spacing,
                line_spacing_rule=line_spacing_rule,
                keep_with_next=True,
                page_break_before=starts_page,
            )
            _set_runs_case(paragraph, cfg.title_text_case)
            touched += 1
            continue

        # Автоматическую нумерацию Word снимаем, если в исходном тексте нет
        # явного маркера в начале абзаца. Раньше это делалось только для пунктов
        # реферата с тире в середине, поэтому строки после двоеточия («Объект
        # работы: …») оставались нумерованными — парсер «додумывал» цифры,
        # которых в исходнике не было. Пункт с уже напечатанным тире наоборот
        # сохраняет маркер, но теряет автобуллет, иначе получалось «- - пункт».
        if not _has_explicit_list_marker(text):
            if _strip_auto_numbering(paragraph):
                stats["auto_numbering_removed"] = (
                    stats.get("auto_numbering_removed", 0) + 1
                )
        elif _strip_auto_numbering(paragraph):
            stats["duplicate_markers_removed"] = (
                stats.get("duplicate_markers_removed", 0) + 1
            )

        # Строки подписей в задании («Задание приняла…», «Руководитель…») —
        # с сохранением табуляции, чтобы черта подписи не уползала влево.
        if kind == "task" and (
            _matches_any(text, preset.title_page.signature_keywords)
            or _paragraph_has_underline(paragraph)
        ):
            _format_signature_paragraph(paragraph, preset.title_page.signatures)
            if paragraph in merged_signatures:
                # Слитная строка — это две колонки: заполнитель между ними не
                # нужен, иначе лидер подписи («черта») растянулся бы на всю ширину.
                sig_cfg = preset.title_page.signatures
                _set_right_tab_stop(
                    paragraph,
                    float(getattr(sig_cfg, "tab_stop_mm", 165.0) or 165.0),
                    "none",
                )
            touched += 1
            continue

        # Строки содержания — своя правая табуляция с точечным заполнителем.
        if kind == "content":
            _format_toc_entry(paragraph, preset, stats)
            touched += 1
            continue

        is_list = bool(
            _LIST_MARKER_PREFIX_RE.match(text) or text[:2] in ("- ", "– ", "* ", "• ")
        )
        _apply_text_block_style(
            paragraph,
            font_name=cfg.body_font_name,
            font_size=cfg.body_font_size,
            color=preset.body.color,
            bold=None,
            italic=None,
            alignment=cfg.body_alignment,
            space_before=0.0,
            space_after=0.0,
            first_line_indent_mm=0.0 if is_list else cfg.body_first_line_indent,
            left_indent_mm=preset.lists.left_indent if is_list else None,
            line_spacing=line_spacing,
            line_spacing_rule=line_spacing_rule,
        )
        touched += 1
    stats[f"special_{kind}_formatted"] = (
        stats.get(f"special_{kind}_formatted", 0) + touched
    )
    return touched


#: Порядок детей ``w:sectPr`` по схеме OOXML (для вставки ``w:pgNumType``).
_SECTPR_TAG_SEQ = (
    "w:footnotePr", "w:endnotePr", "w:type", "w:pgSz", "w:pgMar", "w:paperSrc",
    "w:pgBorders", "w:lnNumType", "w:pgNumType", "w:cols", "w:formProt",
    "w:vAlign", "w:noEndnote", "w:titlePg", "w:textDirection", "w:bidi",
    "w:rtlGutter", "w:docGrid", "w:printerSettings", "w:sectPrChange",
)


def _set_section_page_start(section: Any, start_at: int) -> None:
    """Записывает ``w:pgNumType/@w:start`` — номер первой страницы секции."""
    sectPr = section._sectPr
    element = sectPr.find(qn("w:pgNumType"))
    if element is None:
        element = OxmlElement("w:pgNumType")
        _insert_in_order(sectPr, element, _SECTPR_TAG_SEQ)
    element.set(qn("w:start"), str(int(start_at)))


def _append_page_field(paragraph: Paragraph, field_code: str, cfg: PageNumberingStyle) -> None:
    """Добавляет в абзац поле ``PAGE`` (номер страницы) с оформлением из пресета."""
    run = paragraph.add_run()
    r = run._r
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {field_code} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, placeholder, end):
        r.append(element)
    set_run_font(run, font_name=cfg.font_name, size_pt=cfg.font_size)


def _write_page_number(paragraph: Paragraph, cfg: PageNumberingStyle) -> None:
    """Перезаписывает абзац колонтитула форматом ``page_format`` с полем PAGE."""
    for run in list(paragraph.runs):
        run._r.getparent().remove(run._r)
    for part in re.split(r"(\{PAGE\})", cfg.page_format or "{PAGE}"):
        if part == "{PAGE}":
            _append_page_field(paragraph, "PAGE", cfg)
        elif part:
            run = paragraph.add_run(part)
            set_run_font(run, font_name=cfg.font_name, size_pt=cfg.font_size)
    apply_paragraph_format(
        paragraph.paragraph_format,
        alignment=cfg.alignment,
        space_before=0.0,
        space_after=0.0,
        line_spacing=1.0,
        line_spacing_rule=LineSpacingRule.SINGLE,
        first_line_indent_mm=0.0,
        left_indent_mm=0.0,
    )


def _configure_section_numbering(
    section: Any, cfg: PageNumberingStyle, *, show: bool
) -> None:
    """Настраивает колонтитул секции: с номером страницы либо пустой (скрытый)."""
    container = section.footer if cfg.position == "footer" else section.header
    container.is_linked_to_previous = False
    paragraph = container.paragraphs[0] if container.paragraphs else container.add_paragraph()
    if show:
        _write_page_number(paragraph, cfg)
    else:
        for run in list(paragraph.runs):
            run._r.getparent().remove(run._r)


# --------------------------------------------------------------------------------------
# Основной класс
# --------------------------------------------------------------------------------------


class DocxFormatter:
    """Генерация и нормализация DOCX по пресету :class:`StylePreset`.

    Публичное API (соответствует ТЗ):
      * :meth:`create_document` — собрать новый документ из структуры;
      * :meth:`normalize_document` — вычистить и переформатировать существующий.

    Каждая операция обновляет :attr:`last_stats` — словарь со счётчиками
    (сколько секций/параграфов/таблиц обработано, сколько «костылей» вычищено).
    Он используется обёртками Tool Calling для поля ``details``.
    """

    def __init__(
        self,
        output_dir: str | os.PathLike[str] | None = None,
        log_cb: Any = None,
    ) -> None:
        self.output_dir = Path(output_dir) if output_dir else DEFAULT_OUTPUT_DIR
        self._log_cb = log_cb
        self.last_stats: dict[str, Any] = {}

    # ------------------------------------------------------------------ логирование

    def _log(self, text: str) -> None:
        try:
            if self._log_cb is not None:
                self._log_cb(text)
            else:
                _log(text)
        except Exception:  # noqa: BLE001
            pass

    # ------------------------------------------------------------------ пути

    def _ensure_output_dir(self, path: Path) -> None:
        try:
            path.parent.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            raise DocxFileError(
                f"Не удалось создать каталог {path.parent}: {exc}"
            ) from exc

    def _resolve_create_path(self, output_path: str | None, preset_name: str) -> Path:
        """Куда сохранять создаваемый документ."""
        if output_path:
            candidate = Path(output_path)
            if candidate.suffix.lower() != ".docx":
                candidate = candidate.with_suffix(".docx")
            if not candidate.is_absolute():
                candidate = Path.cwd() / candidate
            return candidate.resolve()
        stem = _safe_stem(preset_name) or "document"
        return (self.output_dir / f"{stem}_{_timestamp()}.docx").resolve()

    def _resolve_normalize_path(self, source: Path, output_path: str, preset_name: str) -> Path:
        """Куда сохранять нормализованный документ (пустой output_path = авто)."""
        if output_path:
            candidate = Path(output_path)
            if candidate.suffix.lower() != ".docx":
                candidate = candidate.with_suffix(".docx")
            if not candidate.is_absolute():
                candidate = Path.cwd() / candidate
            return candidate.resolve()
        stem = _safe_stem(preset_name) or "preset"
        return source.with_name(f"{source.stem}_{stem}.docx").resolve()

    # ------------------------------------------------------------------ публичное API

    def create_document(
        self,
        structure: dict[str, Any] | list[Any],
        preset: StylePreset | str = "gost_standard",
        output_path: str | None = None,
    ) -> str:
        """Создаёт ``.docx`` из структуры контента по пресету.

        Параметры:
          structure — список блоков либо словарь ``{"title": ..., "blocks": [...]}``
            (поддерживается вложенность через ``children``);
          preset — объект :class:`StylePreset`, имя пресета из ``presets/``,
            путь к ``.json`` или inline-JSON;
          output_path — необязательный путь сохранения (иначе в
            ``generated_files/word_automation/``).

        Возвращает путь к сохранённому файлу.
        """
        resolved = resolve_preset(preset)
        blocks = normalize_structure(structure)
        target = self._resolve_create_path(output_path, resolved.name)
        self._ensure_output_dir(target)

        document = docx.Document()
        stats: dict[str, Any] = {
            "preset": resolved.name,
            "blocks": len(blocks),
            "sections": self._apply_page(document, resolved),
            "styles_updated": self._apply_styles(document, resolved),
            "headings": 0,
            "paragraphs": 0,
            "lists": 0,
            "list_items": 0,
            "tables": 0,
            "table_rows": 0,
            "captions": 0,
            "page_breaks": 0,
            "warnings": [],
        }
        section_kinds = self._emit_blocks(document, blocks, resolved, stats)
        # Новые секции (титульный лист/задание) создаются во время рендера —
        # повторно накатываем поля/ориентацию на них.
        stats["sections"] = self._apply_page(document, resolved)
        if resolved.page_numbering.enabled:
            stats["page_numbering"] = self._apply_page_numbering(
                document, resolved, section_kinds
            )
        ensure_document_not_empty(document)

        try:
            document.save(str(target))
        except (OSError, ValueError) as exc:
            raise DocxFileError(
                f"Не удалось сохранить документ {target}: {exc}"
            ) from exc

        stats["file_path"] = str(target)
        stats["size_bytes"] = _file_size(target)
        self.last_stats = stats
        self._log(
            "word_automation: создан документ '%s' (пресет %s, блоков %d)"
            % (target.name, resolved.name, len(blocks))
        )
        return str(target)


    def normalize_document(
        self,
        source_path: str,
        output_path: str = "",
        preset: StylePreset | str = "gost_standard",
        keep_emphasis: bool = True,
    ) -> str:
        """Приводит существующий ``.docx`` к пресету, вычищая «грязное» оформление.

        Выполняется:
          1. очистка текста — серии пробелов/табуляций, неразрывные пробелы,
             «хвостовые» пробелы (с учётом сдвига: подпись -> правое выравнивание);
          2. удаление пустых параграфов-распорок с переводом их в ``space_before``
             следующего смыслового блока;
          3. сброс «сбитых» инлайн-шрифтов и кеглей на значения пресета
             (полужирный/курсив сохраняются при ``keep_emphasis=True``);
          4. принудительный накат пресета на секции, стили ``styles.xml``,
             параграфы и таблицы (включая ``tblHeader``/``cantSplit``).

        Возвращает путь к сохранённому файлу.
        """
        source = Path(source_path)
        if not source.exists():
            raise DocxFileError(f"Исходный документ не найден: {source}")
        if not source.is_file():
            raise DocxFileError(f"Путь не является файлом: {source}")
        if source.suffix.lower() not in (".docx", ".docm"):
            raise DocxFileError(
                f"Ожидается файл .docx, получено '{source.suffix or 'без расширения'}': {source}"
            )

        resolved = resolve_preset(preset)
        target = self._resolve_normalize_path(source, output_path, resolved.name)
        self._ensure_output_dir(target)
        if target.resolve() == source.resolve():
            raise DocxFileError(
                "output_path совпадает с исходным файлом — укажите другой путь, "
                "чтобы не перезаписать оригинал."
            )

        try:
            document = docx.Document(str(source))
        except Exception as exc:  # noqa: BLE001 - python-docx бросает разные типы
            raise DocxFileError(
                f"Не удалось открыть документ {source}: {exc}"
            ) from exc

        stats: dict[str, Any] = {
            "preset": resolved.name,
            "source_path": str(source),
            "sections": self._apply_page(document, resolved),
            "styles_updated": self._apply_styles(document, resolved),
            "paragraphs_total": 0,
            "paragraphs_formatted": 0,
            "headings": 0,
            "captions": 0,
            "lists": 0,
            "empty_paragraphs_removed": 0,
            "space_shifts_fixed": 0,
            "right_aligned_from_shift": 0,
            "multispace_cleaned": 0,
            "runs_rebuilt": 0,
            "tables": 0,
            "table_rows": 0,
            "media_paragraphs_skipped": 0,
            "warnings": [],
        }

        stats["paragraphs_total"] = len(document.paragraphs)
        self._sanitize_body(document, resolved, stats, keep_emphasis)
        stats["tables"], stats["table_rows"] = self._normalize_tables(document, resolved)
        if resolved.page_numbering.enabled:
            stats["page_numbering"] = self._apply_page_numbering(document, resolved)
        ensure_document_not_empty(document)

        try:
            document.save(str(target))
        except (OSError, ValueError) as exc:
            raise DocxFileError(
                f"Не удалось сохранить документ {target}: {exc}"
            ) from exc

        stats["file_path"] = str(target)
        stats["size_bytes"] = _file_size(target)
        self.last_stats = stats
        self._log(
            "word_automation: '%s' -> '%s' (пресет %s; пустых абзацев удалено %d, "
            "сдвигов пробелами исправлено %d)"
            % (
                source.name, target.name, resolved.name,
                stats["empty_paragraphs_removed"], stats["space_shifts_fixed"],
            )
        )
        return str(target)

    # ------------------------------------------------------------------ секции/стили

    def _apply_page(self, document: DocumentObject, preset: StylePreset) -> int:
        """Накатывает поля и ориентацию на все секции документа."""
        page = preset.page
        count = 0
        for section in document.sections:
            section.top_margin = Mm(page.top)
            section.bottom_margin = Mm(page.bottom)
            section.left_margin = Mm(page.left)
            section.right_margin = Mm(page.right)

            width = section.page_width or Mm(A4_SHORT_MM)
            height = section.page_height or Mm(A4_LONG_MM)
            short, long_side = min(width, height), max(width, height)
            if page.is_landscape:
                section.page_width, section.page_height = long_side, short
                section.orientation = WD_ORIENT.LANDSCAPE
            else:
                section.page_width, section.page_height = short, long_side
                section.orientation = WD_ORIENT.PORTRAIT
            count += 1
        return count

    def _apply_styles(self, document: DocumentObject, preset: StylePreset) -> int:
        """Приводит стили ``styles.xml`` к пресету (Normal, Heading N, Caption).

        Это «фундамент» документа: даже параграфы без прямого форматирования
        выглядят по пресету. Прямое форматирование, накатываемое отдельно,
        имеет приоритет над стилями.
        """
        updated = 0
        body = preset.body

        normal = _style_or_none(document, "Normal")
        if normal is not None:
            set_style_font(
                normal,
                font_name=body.font_name,
                size_pt=body.font_size,
                color=body.color,
                bold=False,
                italic=False,
            )
            apply_paragraph_format(
                normal.paragraph_format,
                alignment=body.alignment,
                space_before=body.space_before,
                space_after=body.space_after,
                line_spacing=body.line_spacing,
                line_spacing_rule=body.line_spacing_rule,
                first_line_indent_mm=body.first_line_indent,
                keep_with_next=body.keep_with_next,
                widow_control=body.widow_control,
            )
            updated += 1

        for level in range(1, 10):
            style = _style_or_none(document, f"Heading {level}")
            if style is None:
                continue
            cfg = preset.heading(level)
            set_style_font(
                style,
                font_name=cfg.font_name,
                size_pt=cfg.font_size,
                color=cfg.color,
                bold=cfg.bold,
                italic=cfg.italic,
            )
            apply_paragraph_format(
                style.paragraph_format,
                alignment=cfg.alignment,
                space_before=cfg.space_before,
                space_after=cfg.space_after,
                line_spacing=body.line_spacing,
                line_spacing_rule=body.line_spacing_rule,
                first_line_indent_mm=cfg.first_line_indent,
                keep_with_next=cfg.keep_with_next,
                widow_control=body.widow_control,
                page_break_before=cfg.page_break_before,
            )
            updated += 1

        caption_style = _style_or_none(document, "Caption")
        if caption_style is not None:
            cap = preset.captions
            set_style_font(
                caption_style,
                font_name=cap.font_name,
                size_pt=cap.font_size,
                color=cap.color,
                bold=cap.bold,
                italic=cap.italic,
            )
            apply_paragraph_format(
                caption_style.paragraph_format,
                alignment=cap.alignment,
                space_before=cap.space_before,
                space_after=cap.space_after,
                line_spacing=1.0,
                line_spacing_rule=LineSpacingRule.SINGLE,
                first_line_indent_mm=0.0,
                keep_with_next=False,
                widow_control=body.widow_control,
            )
            updated += 1
        return updated

    # ------------------------------------------------------------------ нумерация страниц

    def _apply_page_numbering(
        self,
        document: DocumentObject,
        preset: StylePreset,
        section_kinds: Sequence[str | None] | None = None,
    ) -> int:
        """Настраивает колонтитул с номером страницы по блоку ``page_numbering``.

        ``section_kinds`` — необязательный список типов секций (``title_page`` /
        ``task`` / ``None``) в порядке ``document.sections``. Он позволяет скрыть
        номер на титульном листе и задании. Нумерация стартует со значения
        ``start_at`` на первой секции и далее идёт непрерывно, поэтому при
        ``start_at=1`` и скрытых титульнике/задании первый видимый номер — 3.
        """
        cfg = preset.page_numbering
        if not cfg.enabled:
            return 0
        sections = list(document.sections)
        if not sections:
            return 0

        _set_section_page_start(sections[0], cfg.start_at)
        kinds = list(section_kinds or [])
        touched = 0
        single_section = len(sections) == 1
        for index, section in enumerate(sections):
            kind = kinds[index] if index < len(kinds) else None
            show = True
            if cfg.hide_on_special:
                if kind == "title_page" and not cfg.show_on_first_page:
                    show = False
                elif kind == "task" and not cfg.show_on_task_page:
                    show = False
            # Документ из одной секции: прячем номер на первой странице (titlePg).
            if single_section and index == 0 and not cfg.show_on_first_page and kind is None:
                section.different_first_page_header_footer = True
            _configure_section_numbering(section, cfg, show=show)
            touched += 1
        return touched

    # ------------------------------------------------------------------ рендер блоков

    def _emit_blocks(
        self,
        document: DocumentObject,
        blocks: list[ContentBlock],
        preset: StylePreset,
        stats: dict[str, Any],
    ) -> list[str | None]:
        """Последовательно добавляет блоки; возвращает типы секций документа.

        Титульный лист при ``title_page.own_section`` выделяется в отдельную
        секцию Word (``w:sectPr``), чтобы далее не наследовать его настройки и
        чтобы можно было скрыть на нём номер страницы.
        """
        section_kinds: list[str | None] = [None]
        for block in blocks:
            try:
                if block.kind == "heading":
                    self._add_heading(document, block, preset)
                    stats["headings"] += 1
                elif block.kind == "paragraph":
                    self._add_paragraph(document, block, preset)
                    stats["paragraphs"] += 1
                elif block.kind == "list":
                    self._add_list(document, block, preset)
                    stats["lists"] += 1
                    stats["list_items"] += len(block.items)
                elif block.kind == "table":
                    self._add_table(document, block, preset)
                    stats["tables"] += 1
                    stats["table_rows"] += len(block.rows)
                elif block.kind == "caption":
                    self._add_caption(document, block, preset)
                    stats["captions"] += 1
                elif block.kind == "page_break":
                    self._add_page_break(document)
                    stats["page_breaks"] += 1
                elif block.kind == "title_page":
                    self._add_title_page(document, block, preset)
                    stats["title_page"] = stats.get("title_page", 0) + 1
                    if preset.title_page.own_section:
                        section_kinds[-1] = "title_page"
                        document.add_section(WD_SECTION.NEW_PAGE)
                        section_kinds.append(None)
                elif block.kind in ("task", "referat", "content"):
                    self._add_special_section(document, block, preset)
                    stats[f"{block.kind}_sections"] = (
                        stats.get(f"{block.kind}_sections", 0) + 1
                    )
                    if (
                        block.kind == "task"
                        and preset.page_numbering.enabled
                        and preset.page_numbering.hide_on_special
                        and not preset.page_numbering.show_on_task_page
                    ):
                        section_kinds[-1] = "task"
                        document.add_section(WD_SECTION.NEW_PAGE)
                        section_kinds.append(None)
                else:
                    note = f"Пропущен блок неизвестного типа: {block.kind!r}"
                    stats["warnings"].append(note)
                    _log(f"word_automation: {note}")
            except DocxStructureError:
                raise
            except Exception as exc:  # noqa: BLE001 - один битый блок не рушит документ
                note = f"Блок ({block.kind}) пропущен: {exc}"
                stats["warnings"].append(note)
                _log(f"word_automation: {note}")
        return section_kinds

    def _add_heading(
        self, document: DocumentObject, block: ContentBlock, preset: StylePreset
    ) -> Paragraph:
        """Добавляет заголовок уровнем 1..9 по пресету."""
        paragraph = document.add_paragraph()
        try:
            paragraph.style = document.styles[f"Heading {block.level}"]
        except (KeyError, ValueError):
            paragraph.style = document.styles["Normal"]
        self._fill_runs(paragraph, block, preset, heading=True)
        cfg = preset.heading(block.level)
        # Структурные разделы (ВВЕДЕНИЕ, ЗАКЛЮЧЕНИЕ, СПИСКИ, ПРИЛОЖЕНИЯ) обязаны
        # начинаться со строго новой страницы — форсируем page_break_before.
        page_break = cfg.page_break_before or _is_structural_title(block.text, preset)
        apply_paragraph_format(
            paragraph.paragraph_format,
            alignment=cfg.alignment,
            space_before=cfg.space_before,
            space_after=cfg.space_after,
            line_spacing=preset.body.line_spacing,
            line_spacing_rule=preset.body.line_spacing_rule,
            first_line_indent_mm=cfg.first_line_indent,
            keep_with_next=cfg.keep_with_next,
            widow_control=preset.body.widow_control,
            page_break_before=page_break,
        )
        return paragraph

    def _add_styled_paragraph(
        self,
        document: DocumentObject,
        text: str,
        *,
        font_name: str,
        font_size: float,
        color: str,
        bold: bool | None,
        italic: bool | None,
        alignment: Alignment,
        space_before: float = 0.0,
        space_after: float = 0.0,
        first_line_indent_mm: float = 0.0,
        left_indent_mm: float | None = None,
        line_spacing: float | None = None,
        line_spacing_rule: LineSpacingRule | None = None,
        case: TextCase | None = None,
        page_break_before: bool | None = None,
        right_tab_stop_mm: float | None = None,
        right_tab_leader: str = "none",
    ) -> Paragraph:
        """Добавляет абзац с явным оформлением (титульный лист, спецразделы)."""
        paragraph = document.add_paragraph()
        if case == TextCase.UPPERCASE:
            text = text.upper()
        run = paragraph.add_run(text)
        set_run_font(
            run,
            font_name=font_name,
            size_pt=font_size,
            color=color,
            bold=bold,
            italic=italic,
        )
        apply_paragraph_format(
            paragraph.paragraph_format,
            alignment=alignment,
            space_before=space_before,
            space_after=space_after,
            line_spacing=line_spacing,
            line_spacing_rule=line_spacing_rule,
            first_line_indent_mm=first_line_indent_mm,
            left_indent_mm=left_indent_mm,
            page_break_before=page_break_before,
        )
        if right_tab_stop_mm is not None:
            _set_right_tab_stop(paragraph, right_tab_stop_mm, right_tab_leader)
        return paragraph

    def _add_title_page(
        self, document: DocumentObject, block: ContentBlock, preset: StylePreset
    ) -> None:
        """Собирает титульный лист: шапка, заголовок работы, подписи, подвал.

        К титульному листу не применяется ``body.first_line_indent`` — все блоки
        выравниваются по своим правилам, подвал прижимается к низу значением
        ``footer.space_before``.
        """
        cfg = preset.title_page
        payload = block.payload or {}
        header_lines = self._payload_lines(payload, ("header", "head", "top", "school"))
        title_text = _as_text(
            _first_present(payload, ("title", "work", "subject", "theme"))
        ).strip()
        signature_lines = self._payload_lines(
            payload, ("signatures", "signature", "sign", "performers", "students")
        )
        footer_lines = self._payload_lines(
            payload, ("footer", "city_year", "city", "bottom", "year")
        )
        if not (header_lines or title_text or signature_lines or footer_lines):
            header_lines = [
                line.strip()
                for line in _as_text(
                    _first_present(payload, ("text", "content", "lines"))
                ).split("\n")
                if line.strip()
            ]

        for line in header_lines:
            b = cfg.header
            self._add_styled_paragraph(
                document, line, font_name=b.font_name, font_size=b.font_size,
                color=b.color, bold=b.bold or None, italic=b.italic or None,
                alignment=b.alignment, space_before=b.space_before,
                space_after=b.space_after, first_line_indent_mm=0.0,
            )
        if title_text:
            b = cfg.title
            self._add_styled_paragraph(
                document, title_text, font_name=b.font_name, font_size=b.font_size,
                color=b.color, bold=b.bold or None, italic=b.italic or None,
                alignment=b.alignment,
                space_before=b.space_before if cfg.distribute_vertically else 0.0,
                space_after=b.space_after, first_line_indent_mm=0.0,
            )
        for line in signature_lines:
            b = cfg.signatures
            has_tab = "\t" in line
            paragraph = self._add_styled_paragraph(
                document, line, font_name=b.font_name, font_size=b.font_size,
                color=b.color, bold=b.bold or None, italic=b.italic or None,
                alignment=Alignment.LEFT if has_tab else b.alignment,
                space_before=b.space_before,
                space_after=b.space_after, first_line_indent_mm=0.0,
                left_indent_mm=None if has_tab else b.left_indent,
            )
            if has_tab and b.preserve_tabs:
                _set_right_tab_stop(paragraph, b.tab_stop_mm, b.tab_leader)
        for line in footer_lines:
            b = cfg.footer
            self._add_styled_paragraph(
                document, line, font_name=b.font_name, font_size=b.font_size,
                color=b.color, bold=b.bold or None, italic=b.italic or None,
                alignment=b.alignment,
                space_before=b.space_before if cfg.distribute_vertically else 0.0,
                space_after=b.space_after, first_line_indent_mm=0.0,
            )

    def _add_special_section(
        self, document: DocumentObject, block: ContentBlock, preset: StylePreset
    ) -> None:
        """Собирает спецраздел: название по центру, текст по ширине, списки as-is."""
        cfg = preset.special_sections
        kind = block.kind
        payload = block.payload or {}
        title = str(payload.get("_title") or "").strip()
        if not title:
            titles = cfg.titles_for(kind)
            title = titles[0].capitalize() if titles else kind.capitalize()
        if cfg.enabled and kind in cfg.apply_to:
            # Бланк задания — компактная форма: одинарный интервал, без раздувания.
            is_blank = kind == "task"
            line_spacing = 1.0 if is_blank else cfg.body_line_spacing
            line_spacing_rule = (
                LineSpacingRule.SINGLE if is_blank else cfg.body_line_spacing_rule
            )
            self._add_styled_paragraph(
                document, title, font_name=cfg.title_font_name,
                font_size=cfg.title_font_size, color=cfg.title_color,
                bold=cfg.title_bold, italic=cfg.title_italic,
                alignment=cfg.title_alignment, space_before=cfg.title_space_before,
                space_after=cfg.title_space_after,
                first_line_indent_mm=cfg.title_first_line_indent,
                line_spacing=line_spacing,
                line_spacing_rule=line_spacing_rule,
                case=cfg.title_text_case,
                page_break_before=cfg.page_break_before,
            )
            for line in payload.get("_body", ()):
                self._add_special_line(document, str(line), kind, preset)
            for item in payload.get("_items", ()):
                self._add_special_line(document, str(item), kind, preset)
            return
        # Спецраздел выключен: простые абзацы по правилам body.
        lines = list(payload.get("_body", ())) + [title] + list(payload.get("_items", ()))
        for line in lines:
            self._add_styled_paragraph(
                document, str(line), font_name=preset.body.font_name,
                font_size=preset.body.font_size, color=preset.body.color,
                bold=None, italic=None, alignment=preset.body.alignment,
                first_line_indent_mm=preset.body.first_line_indent,
                line_spacing=preset.body.line_spacing,
                line_spacing_rule=preset.body.line_spacing_rule,
            )

    def _add_special_line(
        self, document: DocumentObject, text: str, kind: str, preset: StylePreset
    ) -> None:
        """Добавляет строку спецраздела: подпись с табуляцией, TOC-строка или текст.

        Важно: маркер списка НЕ навязывается — если строка уже начинается с тире
        или буквенного/числового перечисления, она выводится как есть.
        """
        cfg = preset.special_sections
        sig_cfg = preset.title_page.signatures
        # 1. Строка подписи (в задании): «Руководитель…», «Задание приняла…».
        is_signature = kind == "task" and (
            _matches_any(text, preset.title_page.signature_keywords)
            or _UNDERLINE_GAP_RE.search(text)
            or "\t" in text
        )
        if is_signature:
            has_tab = "\t" in text
            paragraph = self._add_styled_paragraph(
                document, text, font_name=sig_cfg.font_name,
                font_size=sig_cfg.font_size, color=sig_cfg.color,
                bold=sig_cfg.bold or None, italic=sig_cfg.italic or None,
                alignment=Alignment.LEFT if has_tab else sig_cfg.alignment,
                space_before=sig_cfg.space_before, space_after=sig_cfg.space_after,
                first_line_indent_mm=0.0,
                left_indent_mm=None if has_tab else sig_cfg.left_indent,
            )
            if has_tab:
                _set_right_tab_stop(paragraph, sig_cfg.tab_stop_mm, sig_cfg.tab_leader)
            return
        # 2. Строка содержания: «Заголовок<таб>номер» с точечным заполнителем.
        if kind == "content":
            final_text = text
            if cfg.toc_rebuild_entries:
                match = re.match(r"^(.*?)[\s.\u2026]+\s*(\d{1,4})\s*$", text.strip())
                if match and match.group(1).strip():
                    final_text = (
                        f"{match.group(1).rstrip(' .\t\u2026')}\t{match.group(2)}"
                    )
            self._add_styled_paragraph(
                document, final_text, font_name=cfg.body_font_name,
                font_size=cfg.body_font_size, color=preset.body.color,
                bold=None, italic=None, alignment=cfg.toc_entry_alignment,
                first_line_indent_mm=0.0, left_indent_mm=0.0,
                line_spacing=cfg.body_line_spacing,
                line_spacing_rule=cfg.body_line_spacing_rule,
                right_tab_stop_mm=cfg.toc_right_tab_mm, right_tab_leader=cfg.toc_leader,
            )
            return
        # 3. Обычный текст/пункт — маркер не добавляем.
        is_list = bool(
            _LIST_MARKER_PREFIX_RE.match(text) or text[:2] in ("- ", "– ", "* ", "• ")
        )
        # Бланк задания — одиночный интервал (см. _add_special_section).
        body_line_spacing = 1.0 if kind == "task" else cfg.body_line_spacing
        body_line_rule = LineSpacingRule.SINGLE if kind == "task" else cfg.body_line_spacing_rule
        self._add_styled_paragraph(
            document, text, font_name=cfg.body_font_name,
            font_size=cfg.body_font_size, color=preset.body.color,
            bold=None, italic=None, alignment=cfg.body_alignment,
            first_line_indent_mm=0.0 if is_list else cfg.body_first_line_indent,
            left_indent_mm=preset.lists.left_indent if is_list else None,
            line_spacing=body_line_spacing,
            line_spacing_rule=body_line_rule,
        )

    @staticmethod
    def _payload_lines(payload: dict[str, Any], keys: Sequence[str]) -> list[str]:
        """Извлекает список строк из одного из ключей payload (список или текст)."""
        raw = _first_present(payload, keys)
        if raw is None:
            return []
        if isinstance(raw, (list, tuple)):
            return [str(item).strip() for item in raw if str(item).strip()]
        return [line.strip() for line in _as_text(raw).split("\n") if line.strip()]

    def _add_paragraph(
        self, document: DocumentObject, block: ContentBlock, preset: StylePreset
    ) -> Paragraph:
        """Добавляет обычный абзац по пресету основного текста."""
        paragraph = document.add_paragraph()
        self._fill_runs(paragraph, block, preset, heading=False)
        body = preset.body
        apply_paragraph_format(
            paragraph.paragraph_format,
            alignment=body.alignment,
            space_before=body.space_before,
            space_after=body.space_after,
            line_spacing=body.line_spacing,
            line_spacing_rule=body.line_spacing_rule,
            first_line_indent_mm=body.first_line_indent,
            keep_with_next=body.keep_with_next,
            widow_control=body.widow_control,
        )
        return paragraph

    def _add_page_break(self, document: DocumentObject) -> None:
        """Добавляет принудительный разрыв страницы."""
        paragraph = document.add_paragraph()
        run = paragraph.add_run()
        run.add_break(WD_BREAK.PAGE)

    def _fill_runs(
        self,
        paragraph: Paragraph,
        block: ContentBlock,
        preset: StylePreset,
        *,
        heading: bool,
    ) -> None:
        """Создаёт раны абзаца с нужным шрифтом/кеглем/начертанием.

        Приоритет источника текста: явные ``block.runs`` > inline-разметка в
        ``block.text`` (``**жирный**``/``*курсив*``) > чистый текст. Регистр
        заголовка определяется правилом ``text_case`` пресета.
        """
        if heading:
            cfg = preset.heading(block.level)
            default_bold: bool | None = cfg.bold
            default_italic: bool | None = cfg.italic
        else:
            cfg = preset.body
            default_bold = block.bold
            default_italic = block.italic

        specs = tuple(block.runs) if block.runs else parse_inline_runs(block.text)
        specs = tuple(spec for spec in specs if spec.text)
        if not specs:
            return

        for spec in specs:
            text = cfg.apply_case(spec.text) if heading else spec.text
            if not text:
                continue
            run = paragraph.add_run(text)
            set_run_font(
                run,
                font_name=cfg.font_name,
                size_pt=cfg.font_size,
                color=cfg.color,
                bold=default_bold if spec.bold is None else spec.bold,
                italic=default_italic if spec.italic is None else spec.italic,
                underline=spec.underline,
            )

    def _add_list(
        self, document: DocumentObject, block: ContentBlock, preset: StylePreset
    ) -> None:
        """Добавляет маркированный/нумерованный список (каждый пункт — абзац)."""
        if not block.items:
            return
        body = preset.body
        cfg = preset.lists
        alignment = block.alignment or cfg.alignment

        #: Счётчик нумерации: продолжает авторский номер, чтобы номера не «скакали».
        counter = 0
        for raw_item in block.items:
            text = str(raw_item or "").strip()
            if not text:
                continue
            counter += 1
            paragraph = document.add_paragraph()

            # Исходный маркер снимаем и ставим ровно ОДИН маркер из конфигурации:
            # иначе получалось «- - пункт» или «1. 1. пункт». Авторскую числовую
            # нумерацию («3.», «1.2.») сохраняем, чтобы не сбивать сквозные номера,
            # но синхронизируем с ней счётчик — тогда следующие пункты идут подряд.
            existing = _split_list_marker(text)
            if existing is None:
                marker = cfg.marker_for(ordered=block.ordered, number=counter)
                item_text = text
            else:
                marker_text, rest = existing
                keep_author = block.ordered and bool(
                    _NUMERIC_MARKER_RE.match(marker_text)
                )
                if keep_author:
                    marker = f"{marker_text} "
                    digits = re.match(r"^(\d+)", marker_text)
                    if digits:
                        counter = int(digits.group(1))
                else:
                    marker = cfg.marker_for(ordered=block.ordered, number=counter)
                item_text = rest

            marker_run = paragraph.add_run(marker)
            set_run_font(
                marker_run,
                font_name=body.font_name,
                size_pt=body.font_size,
                color=body.color,
            )

            for spec in parse_inline_runs(item_text):
                if not spec.text:
                    continue
                run = paragraph.add_run(spec.text)
                set_run_font(
                    run,
                    font_name=body.font_name,
                    size_pt=body.font_size,
                    color=body.color,
                    bold=spec.bold,
                    italic=spec.italic,
                    underline=spec.underline,
                )

            apply_paragraph_format(
                paragraph.paragraph_format,
                alignment=alignment,
                space_before=cfg.space_before,
                space_after=cfg.space_after,
                line_spacing=body.line_spacing,
                line_spacing_rule=body.line_spacing_rule,
                first_line_indent_mm=0.0,
                left_indent_mm=cfg.left_indent,
                keep_with_next=cfg.keep_with_next,
                widow_control=body.widow_control,
            )

    def _add_table(
        self, document: DocumentObject, block: ContentBlock, preset: StylePreset
    ) -> Table:
        """Добавляет таблицу с шапкой, границами и форматированием по пресету."""
        cfg = preset.tables
        headers = list(block.header)
        body_rows = [list(row) for row in block.rows]
        columns = _effective_columns(headers, body_rows)
        if columns <= 0:
            raise DocxStructureError("Таблица не имеет ни одной колонки.")

        total_rows = (1 if headers else 0) + max(len(body_rows), 1)
        table = document.add_table(rows=total_rows, cols=columns)
        table.autofit = cfg.autofit
        table.alignment = TABLE_ALIGNMENT_MAP.get(
            block.alignment or cfg.alignment, WD_TABLE_ALIGNMENT.CENTER
        )
        set_table_borders(table, cfg)

        row_index = 0
        if headers:
            self._fill_table_row(table.rows[0], headers, cfg, header=True)
            set_row_properties(
                table.rows[0], cant_split=cfg.cant_split, tbl_header=cfg.repeat_header
            )
            row_index = 1
        for values in body_rows:
            self._fill_table_row(table.rows[row_index], values, cfg, header=False)
            set_row_properties(table.rows[row_index], cant_split=cfg.cant_split)
            row_index += 1

        if block.caption:
            caption_block = ContentBlock(
                kind="caption",
                caption=block.caption,
                caption_kind=block.caption_kind or "table",
                caption_number=block.caption_number,
            )
            caption_paragraph = self._add_caption(document, caption_block, preset)
            # Подпись таблицы по ГОСТ идёт над таблицей — переносим её выше.
            if (
                preset.captions.placement_for(caption_block.caption_kind)
                == CaptionPlacement.ABOVE
            ):
                table._tbl.addprevious(caption_paragraph._p)
        return table

    def _fill_table_row(
        self, row: Any, values: Sequence[str], cfg: TableStyle, *, header: bool
    ) -> None:
        """Заполняет ячейки строки таблицы с форматированием по пресету."""
        alignment = (cfg.header_alignment or cfg.alignment) if header else cfg.alignment
        for index, cell in enumerate(row.cells):
            text = values[index] if index < len(values) else ""
            cell.vertical_alignment = CELL_VALIGN_MAP[cfg.vertical_alignment]
            paragraph = cell.paragraphs[0]
            for stale in list(paragraph.runs):
                stale._r.getparent().remove(stale._r)
            if text:
                run = paragraph.add_run(text)
                set_run_font(
                    run,
                    font_name=cfg.font_name,
                    size_pt=cfg.font_size,
                    color=cfg.color,
                    bold=cfg.header_bold if header else None,
                    italic=cfg.italic or None,
                )
            apply_paragraph_format(
                paragraph.paragraph_format,
                alignment=alignment,
                space_before=cfg.space_before,
                space_after=cfg.space_after,
                line_spacing=cfg.line_spacing,
                line_spacing_rule=cfg.line_spacing_rule,
                first_line_indent_mm=cfg.first_line_indent,
                widow_control=False,
            )

    def _add_caption(
        self, document: DocumentObject, block: ContentBlock, preset: StylePreset
    ) -> Paragraph:
        """Добавляет подпись («Таблица 1 — ...», «Рисунок 2 — ...»)."""
        cap = preset.captions
        paragraph = document.add_paragraph()
        try:
            paragraph.style = document.styles["Caption"]
        except (KeyError, ValueError):
            paragraph.style = document.styles["Normal"]

        text, _placement = _format_caption_text(block, preset)
        if text:
            run = paragraph.add_run(text)
            set_run_font(
                run,
                font_name=cap.font_name,
                size_pt=cap.font_size,
                color=cap.color,
                bold=cap.bold or None,
                italic=cap.italic or None,
            )
        apply_paragraph_format(
            paragraph.paragraph_format,
            alignment=block.alignment or cap.alignment,
            space_before=cap.space_before,
            space_after=cap.space_after,
            line_spacing=1.0,
            line_spacing_rule=LineSpacingRule.SINGLE,
            first_line_indent_mm=0.0,
            keep_with_next=False,
            widow_control=preset.body.widow_control,
        )
        return paragraph

    # ------------------------------------------------------------------ нормализация

    def _sanitize_body(
        self,
        document: DocumentObject,
        preset: StylePreset,
        stats: dict[str, Any],
        keep_emphasis: bool,
    ) -> None:
        """Чистит и форматирует параграфы верхнего уровня (без ячеек таблиц).

        Обход идёт в порядке следования элементов тела документа, чтобы пустые
        абзацы-распорки корректно превращались в ``space_before`` следующего
        смыслового блока (таблица сбрасывает накопленный счётчик).
        """
        body_el = document.element.body
        cleanup = preset.cleanup
        capacity = _line_capacity_chars(preset)
        clean_text = cleanup.text_mode == CleanupMode.CLEAN
        pending_spacers = 0
        spacers: list[Any] = []
        # Распознаём титульный лист и спецразделы, чтобы НЕ накатывать на них
        # параметры body (это и ломало шапку/подписи/город).
        sections = _classify_document_sections(document, preset)
        para_index = -1
        # Внутри раздела «СПИСОК СОКРАЩЕНИЙ И ОБОЗНАЧЕНИЙ» строки вида
        # «АО — акционерное общество» должны оставаться обычными абзацами,
        # поэтому за разделом следим и снимаем с них автоматическую нумерацию.
        in_abbreviations = False

        for child in body_el.iterchildren():
            if child.tag == qn("w:tbl"):
                pending_spacers = 0
                continue
            if child.tag != qn("w:p"):
                continue

            para_index += 1
            paragraph = Paragraph(child, document)
            section_kind = sections.kind_for(para_index)

            if section_kind is not None:
                # Структурный заголовок открывает или закрывает раздел сокращений.
                if section_kind == "structure":
                    in_abbreviations = _matches_any(
                        paragraph.text,
                        ("сокращен", "обозначен", "определения", "термины"),
                    )
                else:
                    in_abbreviations = False
                # Титульный лист, спецразделы и структурные заголовки
                # форматируются отдельными проходами ниже; здесь НЕ накатываем на
                # них параметры body (это ломало шапку/подписи/город). Пустые
                # абзацы-распорки титульника и бланка задания сохраняем, пока не
                # реализована жёсткая табличная сетка.
                is_empty = not paragraph.text.strip() and not element_has_media(paragraph._p)
                drop_spacer = False
                if is_empty:
                    if section_kind == "title_page":
                        drop_spacer = not cleanup.preserve_title_page_spacers
                    elif section_kind in SPECIAL_SECTION_KINDS:
                        drop_spacer = not cleanup.preserve_special_spacers
                if drop_spacer:
                    stats["empty_paragraphs_removed"] += 1
                    remove_paragraph(paragraph)
                pending_spacers = 0
                continue

            raw_text = paragraph.text
            has_media = element_has_media(paragraph._p)
            cleaned = sanitize_text(raw_text, cleanup) if clean_text else raw_text

            if not cleaned.strip() and not has_media:
                # Пустой абзац-распорка (без картинок и полей внутри).
                if cleanup.spacer_mode == SpacerMode.KEEP:
                    continue
                pending_spacers += 1
                stats["empty_paragraphs_removed"] += 1
                if cleanup.spacer_mode == SpacerMode.CONVERT:
                    spacers.append(child)
                else:
                    remove_paragraph(paragraph)
                continue

            if cleaned != raw_text:
                stats["multispace_cleaned"] += 1

            compensated = min(pending_spacers, cleanup.max_spacers)
            extra_before = (
                compensated * cleanup.spacer_spacing
                if cleanup.spacer_mode == SpacerMode.CONVERT and compensated
                else 0.0
            )
            pending_spacers = 0

            shift_alignment = _detect_shift_alignment(raw_text, preset, capacity)
            if shift_alignment is not None:
                stats["space_shifts_fixed"] += 1
                if shift_alignment == Alignment.RIGHT:
                    stats["right_aligned_from_shift"] += 1

            # Автоматическую нумерацию Word снимаем, если в САМОМ исходнике в
            # начале абзаца нет явного маркера. Иначе парсер «додумывал» цифры со
            # скобками там, где их не было (например, в реферате после двоеточия),
            # а строки списка сокращений («АО — акционерное общество») и пункты
            # с уже напечатанным тире получали второй маркер поверх первого.
            if not _has_explicit_list_marker(cleaned):
                if _strip_auto_numbering(paragraph):
                    stats["auto_numbering_removed"] = (
                        stats.get("auto_numbering_removed", 0) + 1
                    )

            kind, cfg = self._paragraph_kind(paragraph, cleaned, preset)
            stats[f"{kind}s"] = stats.get(f"{kind}s", 0) + 1

            self._rewrite_runs(paragraph, cleaned, cleanup, keep_emphasis, has_media, stats)
            self._apply_block_format(paragraph, kind, cfg, preset, extra_before, shift_alignment)
            stats["paragraphs_formatted"] += 1

        # Удаляем накопленные распорки после прохода (порядок элементов уже не важен).
        for element in spacers:
            parent = element.getparent()
            if parent is not None:
                parent.remove(element)

        # Титульный лист и спецразделы — отдельные проходы форматирования.
        if sections.title_page:
            _format_title_page(sections.title_page, preset, stats, document)
            self._break_after_title_page(sections.title_page, preset, stats)
        for _paragraph in sections.structural:
            _format_structure_heading(_paragraph, preset, stats)
        for _level, _paragraph in sections.subheadings:
            _format_subheading(_paragraph, _level, preset, stats)
        for kind, paragraphs in sections.special.items():
            _format_special_section(kind, paragraphs, preset, stats)

    @staticmethod
    def _break_after_title_page(
        title_paragraphs: list[Paragraph], preset: StylePreset, stats: dict[str, Any]
    ) -> None:
        """Ставит разрыв страницы строго ПОСЛЕ последнего абзаца титула.

        Внутри титула все явные разрывы сняты (:func:`_format_title_page`),
        поэтому разрыв назначается следующему абзацу — им оказывается первый
        абзац спецраздела или структурного заголовка. Если следующий абзац уже
        начинает новую страницу, ничего не делаем.
        """
        if not title_paragraphs or not preset.special_sections.page_break_before:
            return
        last = title_paragraphs[-1]._p
        following = last.getnext()
        while following is not None and following.tag != qn("w:p"):
            following = following.getnext()
        if following is None:
            return
        paragraph = Paragraph(following, last.getroottree())
        if _has_explicit_page_break(paragraph):
            return
        paragraph.paragraph_format.page_break_before = True
        stats["title_page_break_after"] = stats.get("title_page_break_after", 0) + 1

    def _paragraph_kind(
        self, paragraph: Paragraph, text: str, preset: StylePreset
    ) -> tuple[str, Any]:
        """Определяет тип абзаца и соответствующий ему конфиг пресета.

        Приоритет: явный стиль (``Heading N`` / ``Caption`` / ``List …``),
        затем эвристика по тексту (маркер списка, шаблон подписи).
        """
        style_name = ""
        try:
            style_name = str(paragraph.style.name or "")
        except (AttributeError, KeyError):  # noqa: PERF203
            style_name = ""
        lowered = style_name.lower()

        if lowered.startswith("heading"):
            digits = "".join(ch for ch in style_name if ch.isdigit())
            level = int(digits) if digits else 1
            return "heading", preset.heading(max(1, min(level, 9)))
        if lowered == "title" or lowered == "subtitle":
            return "heading", preset.heading(1)
        if lowered == "caption":
            return "caption", preset.captions
        if lowered.startswith("list"):
            return "list", preset.lists
        if _CAPTION_RE.match(text.strip()):
            return "caption", preset.captions
        stripped = text.strip()
        if stripped[:2] in ("- ", "– ", "* ", "• ") or _NUMBERED_LEAD_RE.match(stripped):
            return "list", preset.lists
        return "paragraph", preset.body

    def _rewrite_runs(
        self,
        paragraph: Paragraph,
        cleaned: str,
        cleanup: CleanupPolicy,
        keep_emphasis: bool,
        has_media: bool,
        stats: dict[str, Any],
    ) -> None:
        """Пересобирает раны абзаца с очищенным текстом.

        Если все раны имеют одинаковое начертание — схлопываем их в один ран
        (типичный «мусор» после конвертаций из PDF/HTML). При разнородном
        начертании и ``keep_emphasis=True`` чистим каждый ран отдельно, чтобы
        не потерять полужирный/курсив.
        """
        runs = paragraph.runs

        if has_media:
            # Внутри картинки/поля: пересборка запрещена, чистим только текст.
            for run in runs:
                if element_has_media(run._r):
                    continue
                new_text = sanitize_text(run.text, cleanup)
                if new_text != run.text:
                    run.text = new_text
            return

        if not runs:
            if cleaned:
                paragraph.add_run(cleaned)
                stats["runs_rebuilt"] += 1
            return

        if not keep_emphasis or len(runs) == 1:
            runs[0].text = cleaned
            for run in runs[1:]:
                parent = run._r.getparent()
                if parent is not None:
                    parent.remove(run._r)
            stats["runs_rebuilt"] += 1
            return

        emphasis = {(bool(r.bold), bool(r.italic), bool(r.underline)) for r in runs}
        if len(emphasis) == 1:
            runs[0].text = cleaned
            for run in runs[1:]:
                parent = run._r.getparent()
                if parent is not None:
                    parent.remove(run._r)
            stats["runs_rebuilt"] += 1
            return

        for run in runs:
            new_text = sanitize_text(run.text, cleanup)
            if new_text != run.text:
                run.text = new_text
                stats["multispace_cleaned"] += 1

    def _apply_block_format(
        self,
        paragraph: Paragraph,
        kind: str,
        cfg: Any,
        preset: StylePreset,
        extra_before: float,
        shift_alignment: Alignment | None,
    ) -> None:
        """Накатывает пресет на один абзац с учётом «распорочного» отступа."""
        body = preset.body
        if kind == "heading":
            line_spacing, line_rule = body.line_spacing, body.line_spacing_rule
            first_line = cfg.first_line_indent
            alignment = cfg.alignment
            keep_next = cfg.keep_with_next
            page_break = cfg.page_break_before
        elif kind == "caption":
            line_spacing, line_rule = 1.0, LineSpacingRule.SINGLE
            first_line = 0.0
            alignment = cfg.alignment
            keep_next = False
            page_break = False
        elif kind == "list":
            line_spacing, line_rule = body.line_spacing, body.line_spacing_rule
            first_line = 0.0
            alignment = cfg.alignment
            keep_next = cfg.keep_with_next
            page_break = False
        else:
            line_spacing, line_rule = body.line_spacing, body.line_spacing_rule
            first_line = body.first_line_indent
            alignment = body.alignment
            keep_next = body.keep_with_next
            page_break = False

        if shift_alignment is not None:
            # «Выравнивание пробелами» заменено настоящим выравниванием абзаца.
            alignment = shift_alignment
            first_line = 0.0

        left_indent = cfg.left_indent if kind == "list" else 0.0
        apply_paragraph_format(
            paragraph.paragraph_format,
            alignment=alignment,
            space_before=cfg.space_before + extra_before,
            space_after=cfg.space_after,
            line_spacing=line_spacing,
            line_spacing_rule=line_rule,
            first_line_indent_mm=first_line,
            left_indent_mm=left_indent,
            keep_with_next=keep_next,
            widow_control=body.widow_control,
            page_break_before=page_break,
        )

        for run in paragraph.runs:
            if element_has_media(run._r):
                continue
            if kind == "heading":
                set_run_font(
                    run, font_name=cfg.font_name, size_pt=cfg.font_size,
                    color=cfg.color, bold=cfg.bold, italic=cfg.italic,
                )
            elif kind == "caption":
                set_run_font(
                    run, font_name=cfg.font_name, size_pt=cfg.font_size,
                    color=cfg.color, bold=cfg.bold, italic=cfg.italic,
                )
            elif kind == "list":
                set_run_font(
                    run, font_name=body.font_name, size_pt=body.font_size,
                    color=body.color,
                )
            else:
                set_run_font(
                    run, font_name=body.font_name, size_pt=body.font_size,
                    color=body.color,
                )

    def _normalize_tables(
        self, document: DocumentObject, preset: StylePreset
    ) -> tuple[int, int]:
        """Приводит все таблицы документа к пресету (включая ячейки)."""
        cfg = preset.tables
        cleanup = preset.cleanup
        clean_text = cleanup.text_mode == CleanupMode.CLEAN
        tables = 0
        rows = 0
        for table in document.tables:
            tables += 1
            table.autofit = cfg.autofit
            table.alignment = TABLE_ALIGNMENT_MAP.get(
                cfg.alignment, WD_TABLE_ALIGNMENT.CENTER
            )
            set_table_borders(table, cfg)

            for row_index, row in enumerate(table.rows):
                rows += 1
                is_header = row_index == 0 and cfg.repeat_header
                set_row_properties(
                    row, cant_split=cfg.cant_split, tbl_header=True if is_header else None
                )
                alignment = (
                    (cfg.header_alignment or cfg.alignment) if is_header else cfg.alignment
                )
                for cell in row.cells:
                    cell.vertical_alignment = CELL_VALIGN_MAP[cfg.vertical_alignment]
                    for paragraph in cell.paragraphs:
                        if clean_text:
                            for run in paragraph.runs:
                                if element_has_media(run._r):
                                    continue
                                new_text = sanitize_text(run.text, cleanup)
                                if new_text != run.text:
                                    run.text = new_text
                        apply_paragraph_format(
                            paragraph.paragraph_format,
                            alignment=alignment,
                            space_before=cfg.space_before,
                            space_after=cfg.space_after,
                            line_spacing=cfg.line_spacing,
                            line_spacing_rule=cfg.line_spacing_rule,
                            first_line_indent_mm=cfg.first_line_indent,
                            widow_control=False,
                        )
                        for run in paragraph.runs:
                            if element_has_media(run._r):
                                continue
                            set_run_font(
                                run,
                                font_name=cfg.font_name,
                                size_pt=cfg.font_size,
                                color=cfg.color,
                                bold=cfg.header_bold if is_header else None,
                                italic=cfg.italic or None,
                            )
        return tables, rows









# --------------------------------------------------------------------------------------
# Служебные функции
# --------------------------------------------------------------------------------------


def _timestamp() -> str:
    """Метка времени для уникальных имён файлов (YYYYmmdd_HHMMSS)."""
    from datetime import datetime

    return datetime.now().strftime("%Y%m%d_%H%M%S")


def _safe_stem(name: str) -> str:
    """Имя файла без недопустимых символов."""
    text = str(name or "").strip().replace("\\", "/").split("/")[-1]
    text = re.sub(r"[\x00-\x1f<>:\"\\|?*]+", "_", text)
    text = re.sub(r"\s+", "_", text).strip("._")
    if text.lower().endswith(".docx"):
        text = text[:-5]
    return text[:80].strip("._")


def _file_size(path: Path) -> int:
    """Размер файла в байтах (0, если stat недоступен)."""
    try:
        return int(path.stat().st_size)
    except OSError:
        return 0
